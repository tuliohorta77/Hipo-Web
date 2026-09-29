"""
HIPO — Regras puras do RPeR (services/rper.py, rper_ia.py, rper_render.py).

Sem banco e sem rede: rodam no pytest local do Windows. O que so aparece
com banco (a coleta, as rotas, as metas) esta em test_rper.py.
"""
import os
from datetime import datetime, timezone
from io import BytesIO
from uuid import uuid4

import pytest

os.environ.setdefault("DATABASE_URL", "postgresql://hipo_test:hipo_test@localhost:5432/hipo_test")
os.environ.setdefault("JWT_SECRET", "test-secret-key-hipo-2026")

from services import rper as regras  # noqa: E402
from services import rper_ia  # noqa: E402

AGORA = datetime(2026, 9, 1, 12, tzinfo=timezone.utc)
INICIO = datetime(2026, 8, 10, 13, tzinfo=timezone.utc)


def _vazio():
    return {k: [] for k in ("reunioes", "agendamentos", "tarefas", "leads", "vendas",
                            "propostas", "ativas", "parceiros", "parcerias",
                            "indicacoes")} | {"atrasadas": {}}


def _pessoas():
    return {
        "EC": [{"id": uuid4(), "nome": "Aline Martins"}],
        "SDR": [{"id": uuid4(), "nome": "Kethlleen Gomes"}, {"id": uuid4(), "nome": "Gabriel Lira"}],
        "EV": [{"id": uuid4(), "nome": "Jakeline Santana"}, {"id": uuid4(), "nome": "Bruno Gonçalo"}],
    }


SEM_METAS = {"squad": {}, "pessoa": {}}


def _montar(dados, pessoas, metas_fechado=SEM_METAS, metas_novo=SEM_METAS):
    return regras.montar(ano=2026, mes=8, dados=dados, pessoas=pessoas,
                         metas_fechado=metas_fechado, metas_novo=metas_novo, agora=AGORA)


def _linha(bloco, chave):
    return next(l for l in bloco if l["chave"] == chave)


# ── Formatação e meses ───────────────────────────────────────────────

class TestFormatacao:
    @pytest.mark.parametrize("valor,formato,esperado", [
        (18908, "moeda", "R$ 18.908"),
        (226.5, "moeda", "R$ 226,50"),
        (0, "moeda", "R$ 0"),
        (93.75, "percentual", "94%"),
        (33.333, "percentual", "33,3%"),
        (7.0, "percentual", "7%"),
        (1234, "inteiro", "1.234"),
        (None, "inteiro", "—"),
    ])
    def test_formatos_do_slide(self, valor, formato, esperado):
        assert regras.formatar(valor, formato) == esperado

    def test_virada_de_ano(self):
        assert regras.mes_seguinte(2026, 12) == (2027, 1)
        assert regras.mes_anterior(2027, 1) == (2026, 12)

    def test_mes_invalido(self):
        with pytest.raises(regras.RperInvalido):
            regras.janela_do_mes(2026, 13)

    def test_primeiro_nome(self):
        assert regras.primeiro_nome("  bruno gonçalo ") == "BRUNO"


class TestCatalogo:
    def test_tres_squads_com_duas_metas_principais(self):
        """A tabela de meta por pessoa do planejamento tem duas colunas."""
        for squad in regras.SQUADS:
            principais = [i for i in regras.INDICADORES[squad] if i.principal]
            assert len(principais) == 2, squad

    def test_chaves_unicas_por_squad(self):
        for squad, lista in regras.INDICADORES.items():
            chaves = [i.chave for i in lista]
            assert len(chaves) == len(set(chaves)), squad

    def test_indicador_de_outro_squad_e_recusado(self):
        with pytest.raises(regras.RperInvalido):
            regras.validar_indicador("SDR", "pipeline")
        with pytest.raises(regras.RperInvalido):
            regras.validar_squad("Gerente")


# ── As contas ────────────────────────────────────────────────────────

class TestSquadNaoESoma:
    def test_venda_com_dois_evs_conta_uma_vez_no_squad(self):
        """
        O motivo de a coleta devolver linhas: a venda com os dois EVs
        envolvidos e de cada um E e UMA do squad. Somar as pessoas daria 2.
        """
        p = _pessoas()
        ev1, ev2 = (x["id"] for x in p["EV"])
        dados = _vazio()
        dados["vendas"] = [{"id": uuid4(), "valor": 500, "envolvidos": [ev1, ev2]}]
        r = _montar(dados, p)
        ev = r["squads"]["EV"]
        assert _linha(ev["total"], "vendas")["realizado"] == 1
        assert _linha(ev["total"], "nmrr")["realizado"] == 500
        assert [_linha(x["indicadores"], "vendas")["realizado"] for x in ev["pessoas"]] == [1, 1]

    def test_venda_de_quem_nao_e_do_squad_nao_entra(self):
        p = _pessoas()
        dados = _vazio()
        dados["vendas"] = [{"id": uuid4(), "valor": 500, "envolvidos": [uuid4()]}]
        assert _linha(_montar(dados, p)["squads"]["EV"]["total"], "vendas")["realizado"] == 0


class TestSdr:
    def test_agendamento_e_noshow_de_quem_agendou(self):
        p = _pessoas()
        k, g = (x["id"] for x in p["SDR"])
        ev = p["EV"][0]["id"]
        opp = uuid4()
        dados = _vazio()
        dados["agendamentos"] = [
            {"id": uuid4(), "agendado_por": k, "oportunidade_id": opp, "valor": 300},
            {"id": uuid4(), "agendado_por": k, "oportunidade_id": opp, "valor": 300},
            {"id": uuid4(), "agendado_por": g, "oportunidade_id": uuid4(), "valor": None},
        ]
        base = {"anfitriao_id": ev, "parceiro_id": None, "oportunidade_id": opp, "inicio": INICIO}
        dados["reunioes"] = [
            {"id": uuid4(), "agendado_por": k, "efetivo": "realizada", **base},
            {"id": uuid4(), "agendado_por": k, "efetivo": "no_show", **base},
            {"id": uuid4(), "agendado_por": k, "efetivo": None, **base},
            # Reuniao de parceiro e ilha: nao entra no SDR.
            {"id": uuid4(), "agendado_por": k, "efetivo": "no_show",
             **{**base, "parceiro_id": uuid4()}},
        ]
        sdr = _montar(dados, p)["squads"]["SDR"]
        kt = sdr["pessoas"][0]["indicadores"]
        assert _linha(kt, "agendamentos")["realizado"] == 2
        assert _linha(kt, "reunioes_realizadas")["realizado"] == 1
        assert _linha(kt, "noshow")["realizado"] == 50.0
        # Ticket gerado: a MESMA oportunidade agendada duas vezes conta uma.
        assert _linha(kt, "pipeline_gerado")["realizado"] == 300
        assert _linha(sdr["total"], "agendamentos")["realizado"] == 3

    def test_noshow_sem_reuniao_fechada_e_indefinido(self):
        sdr = _montar(_vazio(), _pessoas())["squads"]["SDR"]
        assert _linha(sdr["total"], "noshow")["realizado"] is None
        assert _linha(sdr["total"], "noshow")["realizado_txt"] == "—"

    def test_tarefas_e_contas_prospectadas(self):
        p = _pessoas()
        k = p["SDR"][0]["id"]
        conta = uuid4()
        dados = _vazio()
        dados["tarefas"] = [
            {"id": uuid4(), "responsavel_id": k, "oportunidade_id": uuid4(),
             "conta_id": conta, "parceiro_id": None, "concluida": True},
            {"id": uuid4(), "responsavel_id": k, "oportunidade_id": uuid4(),
             "conta_id": conta, "parceiro_id": None, "concluida": True},
            {"id": uuid4(), "responsavel_id": k, "oportunidade_id": uuid4(),
             "conta_id": uuid4(), "parceiro_id": None, "concluida": False},
            # Tarefa de parceiro nao e prospeccao.
            {"id": uuid4(), "responsavel_id": k, "oportunidade_id": None,
             "conta_id": None, "parceiro_id": uuid4(), "concluida": True},
        ]
        t = _montar(dados, p)["squads"]["SDR"]["pessoas"][0]["indicadores"]
        assert _linha(t, "tarefas")["realizado"] == 3
        assert _linha(t, "contas")["realizado"] == 1
        assert _linha(t, "taxa_execucao")["realizado"] == 66.7


class TestEv:
    def test_pipeline_so_apresentacao_e_negociacao(self):
        p = _pessoas()
        j = p["EV"][0]["id"]
        dados = _vazio()
        dados["ativas"] = [
            {"id": uuid4(), "numero": "1", "fase": "negociacao", "valor": 1000,
             "empresa": "A", "envolvidos": [j], "ultimo_fup": None},
            {"id": uuid4(), "numero": "2", "fase": "apresentacao", "valor": 500,
             "empresa": "B", "envolvidos": [j], "ultimo_fup": None},
            {"id": uuid4(), "numero": "3", "fase": "lead", "valor": 9999,
             "empresa": "C", "envolvidos": [j], "ultimo_fup": None},
        ]
        r = _montar(dados, p)
        ev = r["squads"]["EV"]
        assert _linha(ev["total"], "pipeline")["realizado"] == 1500
        assert _linha(ev["total"], "em_negociacao")["realizado"] == 1
        assert ev["pessoas"][0]["concentracao_pct"] == 67
        assert [o["empresa"] for o in ev["negociacoes"]["top"]] == ["A"]
        assert ev["negociacoes"]["dois_maiores_pct"] == 67
        assert ev["graficos"]["outras_fases"] == [2, 0]

    def test_conversao_e_ticket(self):
        p = _pessoas()
        j = p["EV"][0]["id"]
        dados = _vazio()
        base = {"agendado_por": None, "anfitriao_id": j, "parceiro_id": None,
                "oportunidade_id": uuid4(), "inicio": INICIO}
        dados["reunioes"] = [{"id": uuid4(), "efetivo": "realizada", **base} for _ in range(4)]
        dados["vendas"] = [{"id": uuid4(), "valor": 300, "envolvidos": [j]},
                           {"id": uuid4(), "valor": 100, "envolvidos": [j]}]
        t = _montar(dados, p)["squads"]["EV"]["total"]
        assert _linha(t, "taxa_conversao")["realizado"] == 50.0
        assert _linha(t, "ticket_medio")["realizado"] == 200.0
        assert _linha(t, "ticket_medio")["realizado_txt"] == "R$ 200"

    def test_top_limita_em_dez_e_ordena_por_valor(self):
        p = _pessoas()
        j = p["EV"][0]["id"]
        dados = _vazio()
        dados["ativas"] = [
            {"id": uuid4(), "numero": str(v), "fase": "negociacao", "valor": v,
             "empresa": f"E{v}", "envolvidos": [j], "ultimo_fup": None}
            for v in range(100, 1300, 100)
        ]
        top = _montar(dados, p)["squads"]["EV"]["negociacoes"]["top"]
        assert len(top) == 10
        assert top[0]["valor"] == 1200 and top[-1]["valor"] == 300


class TestEc:
    def test_mrr_e_o_fechado_com_ec_envolvido(self):
        """Decisao do Tulio (29/09): MRR do EC = vendas do mes com ele envolvido."""
        p = _pessoas()
        a = p["EC"][0]["id"]
        dados = _vazio()
        dados["vendas"] = [{"id": uuid4(), "valor": 700, "envolvidos": [a]},
                           {"id": uuid4(), "valor": 999, "envolvidos": [uuid4()]}]
        dados["parceiros"] = [{"id": uuid4(), "ec_id": a} for _ in range(3)]
        conta = uuid4()
        dados["parcerias"] = [{"id": 1, "conta_id": conta, "para_usuario_id": a},
                              {"id": 2, "conta_id": conta, "para_usuario_id": a}]
        t = _montar(dados, p)["squads"]["EC"]["total"]
        assert _linha(t, "mrr")["realizado"] == 700
        assert _linha(t, "contas_gestao")["realizado"] == 3
        assert _linha(t, "parcerias")["realizado"] == 1


class TestMetas:
    def test_meta_do_squad_nao_e_a_soma_das_pessoas(self):
        p = _pessoas()
        k, g = (str(x["id"]) for x in p["SDR"])
        metas = {"squad": {("SDR", "agendamentos"): 70.0},
                 "pessoa": {("SDR", k, "agendamentos"): 40.0, ("SDR", g, "agendamentos"): 40.0}}
        sdr = _montar(_vazio(), p, metas_fechado=metas)["squads"]["SDR"]
        assert _linha(sdr["total"], "agendamentos")["meta"] == 70.0
        assert _linha(sdr["pessoas"][0]["indicadores"], "agendamentos")["meta"] == 40.0

    def test_atingimento_e_carinha(self):
        p = _pessoas()
        dados = _vazio()
        dados["agendamentos"] = [{"id": uuid4(), "agendado_por": p["SDR"][0]["id"],
                                  "oportunidade_id": None, "valor": None} for _ in range(8)]
        metas = {"squad": {("SDR", "agendamentos"): 10.0}, "pessoa": {}}
        l = _linha(_montar(dados, p, metas_fechado=metas)["squads"]["SDR"]["total"], "agendamentos")
        assert l["atingimento_txt"] == "80%"
        assert l["carinha"] == "neutro"

    def test_sem_meta_celulas_vazias(self):
        l = _linha(_montar(_vazio(), _pessoas())["squads"]["EV"]["total"], "nmrr")
        assert l["meta_txt"] == "" and l["atingimento_txt"] == ""

    def test_planejamento_usa_as_metas_do_mes_novo(self):
        p = _pessoas()
        j = str(p["EV"][0]["id"])
        novo = {"squad": {("EV", "nmrr"): 8000.0}, "pessoa": {("EV", j, "nmrr"): 5000.0}}
        pl = _montar(_vazio(), p, metas_novo=novo)["squads"]["EV"]["planejamento"]
        assert {"rotulo": "NMRR", "meta_txt": "R$ 8.000"} in pl["time"]
        assert pl["colunas"] == ["NMRR", "PIPELINE"]
        assert pl["pessoas"][0] == {"rotulo": "JAKELINE", "metas": ["R$ 5.000", ""]}
        # Indicador de posicao sem meta nao ocupa linha no planejamento.
        assert not any(t["rotulo"] == "EM NEGOCIAÇÃO" for t in pl["time"])


# ── Textos ───────────────────────────────────────────────────────────

def _rper_com_dados():
    p = _pessoas()
    j, b = (x["id"] for x in p["EV"])
    dados = _vazio()
    dados["ativas"] = [
        {"id": uuid4(), "numero": "1", "fase": "negociacao", "valor": 11250,
         "empresa": "LIFE RH", "envolvidos": [j], "ultimo_fup": None},
        {"id": uuid4(), "numero": "2", "fase": "negociacao", "valor": 930,
         "empresa": "ESTEIO", "envolvidos": [b], "ultimo_fup": None},
        {"id": uuid4(), "numero": "3", "fase": "negociacao", "valor": None,
         "empresa": "SEM VALOR", "envolvidos": [b], "ultimo_fup": None},
    ]
    dados["atrasadas"] = {j: 4}
    return _montar(dados, p), p


class TestTextosPadrao:
    def test_todos_os_campos_existem(self):
        r, p = _rper_com_dados()
        t = regras.textos_padrao(r)
        for squad in regras.SQUADS:
            assert t[squad]["leitura"]
            assert t[squad]["acoes"]
            assert set(t[squad]["pessoas"]) == {str(x["id"]) for x in p[squad]}
        assert "LIFE RH e ESTEIO" in t["EV"]["foco"]
        assert "Zerar as 4 tarefas atrasadas em aberto" in t["EV"]["acoes"]
        assert any("1 oportunidades sem valor" in a for a in t["EV"]["acoes"])

    def test_texto_padrao_passa_na_propria_guarda(self):
        """O padrao so pode citar numero que esta nos dados."""
        from services.validacao_numerica import numeros_invalidos, numeros_permitidos
        r, _ = _rper_com_dados()
        permitidos = numeros_permitidos(rper_ia.payload_para_ia(r))
        t = regras.textos_padrao(r)
        for squad, item in t.items():
            textos = [item["leitura"], *item["pessoas"].values(), *item["acoes"]]
            textos += [item.get("foco", ""), item.get("pipeline", "")]
            for texto in textos:
                assert numeros_invalidos(texto, permitidos) == [], (squad, texto)


class TestIa:
    def test_sem_resposta_fica_o_padrao(self):
        r, _ = _rper_com_dados()
        padrao = regras.textos_padrao(r)
        textos, stats = rper_ia.mesclar(r, padrao, None)
        assert textos == padrao
        assert stats == {"ia": False, "aceitos": 0, "descartados": 0}

    def test_campo_valido_entra_e_invalido_cai_no_padrao(self):
        r, p = _rper_com_dados()
        padrao = regras.textos_padrao(r)
        jake = p["EV"][0]
        resposta = {
            "EV": {
                "leitura": "Pipeline de R$ 12.180 concentrado em LIFE RH.",
                # 37 nao existe nos dados: so ESTE campo volta ao padrao.
                "pessoas": {jake["nome"]: "Fechou 37 contratos."},
                "acoes": ["Priorizar LIFE RH", "Zerar as 4 tarefas atrasadas"],
                "foco": "LIFE RH e ESTEIO pesam 55% do pipeline.",
            },
        }
        textos, stats = rper_ia.mesclar(r, padrao, resposta)
        assert textos["EV"]["leitura"] == "Pipeline de R$ 12.180 concentrado em LIFE RH."
        assert textos["EV"]["pessoas"][str(jake["id"])] == padrao["EV"]["pessoas"][str(jake["id"])]
        assert textos["EV"]["acoes"] == ["Priorizar LIFE RH", "Zerar as 4 tarefas atrasadas"]
        assert textos["EV"]["foco"] == padrao["EV"]["foco"]
        assert stats["descartados"] == 2
        # Squad que a IA nao mandou fica inteiro no padrao.
        assert textos["SDR"] == padrao["SDR"]

    def test_texto_longo_demais_cai_no_padrao(self):
        r, _ = _rper_com_dados()
        padrao = regras.textos_padrao(r)
        textos, _ = rper_ia.mesclar(r, padrao, {"EC": {"leitura": "a" * 400}})
        assert textos["EC"]["leitura"] == padrao["EC"]["leitura"]

    def test_acoes_so_trocam_com_duas_validas(self):
        r, _ = _rper_com_dados()
        padrao = regras.textos_padrao(r)
        textos, _ = rper_ia.mesclar(r, padrao, {"EV": {"acoes": ["Priorizar LIFE RH", "Ligar 99 vezes"]}})
        assert textos["EV"]["acoes"] == padrao["EV"]["acoes"]

    def test_payload_sem_uuid(self):
        """UUID no JSON liberaria digitos soltos na guarda numerica."""
        r, p = _rper_com_dados()
        texto = str(rper_ia.payload_para_ia(r))
        for squad in p.values():
            for pessoa in squad:
                assert str(pessoa["id"]) not in texto

    @pytest.mark.parametrize("bruto,esperado", [
        ('{"EV": {}}', {"EV": {}}),
        ('```json\n{"EV": {}}\n```', {"EV": {}}),
        ("nao e json", None),
        ("[1, 2]", None),
    ])
    def test_json_da_resposta(self, bruto, esperado):
        assert rper_ia._json_da_resposta(bruto) == esperado

    async def test_sem_chave_nao_chama_a_api(self, monkeypatch):
        from services import ia as ia_base
        monkeypatch.setattr(ia_base, "configurada", lambda: False)
        r, _ = _rper_com_dados()
        padrao = regras.textos_padrao(r)
        textos, stats = await rper_ia.escrever(r, padrao)
        assert textos == padrao and stats["ia"] is False


# ── O arquivo ────────────────────────────────────────────────────────

class TestRender:
    def _abrir(self, corpo):
        from pptx import Presentation
        return Presentation(BytesIO(corpo))

    def test_quinze_slides_na_ordem_da_reuniao(self):
        from services import rper_render
        r, _ = _rper_com_dados()
        prs = self._abrir(rper_render.montar_pptx(r, regras.textos_padrao(r)))
        assert len(prs.slides) == 15
        textos = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
                  for s in prs.slides]
        assert "SETEMBRO | 2026" in textos[0]
        assert "Resultados de agosto/2026" in textos[0]
        assert "EC — EXECUTIVO DE CONTAS" in textos[1]
        assert "SDR — RESULTADOS" in textos[6]
        assert "PIPELINE POR" in textos[12]
        assert "PLANEJAMENTO SETEMBRO" in textos[14]

    def test_tabela_por_pessoa_tem_uma_coluna_por_pessoa_mais_o_squad(self):
        from services import rper_render
        r, _ = _rper_com_dados()
        prs = self._abrir(rper_render.montar_pptx(r, regras.textos_padrao(r)))
        tabela = next(sh.table for sh in prs.slides[11].shapes if sh.has_table)
        cab = [c.text for c in tabela.rows[0].cells]
        assert cab == ["INDICADOR", "JAKELINE", "BRUNO", "SQUAD"]

    def test_squad_sem_ninguem_ainda_gera(self):
        from services import rper_render
        p = _pessoas()
        p["EC"] = []
        r = _montar(_vazio(), p)
        prs = self._abrir(rper_render.montar_pptx(r, regras.textos_padrao(r)))
        assert len(prs.slides) == 15

    def test_nome_do_arquivo(self):
        from services import rper_render
        r, _ = _rper_com_dados()
        assert rper_render.nome_do_arquivo(r, "pptx") == "RPeR_SETEMBRO_2026_CONTROLLER_MEDSEG.pptx"
        r2 = regras.montar(ano=2026, mes=2, dados=_vazio(), pessoas=_pessoas(),
                           metas_fechado=SEM_METAS, metas_novo=SEM_METAS, agora=AGORA)
        assert rper_render.nome_do_arquivo(r2, "pdf") == "RPeR_MARCO_2026_CONTROLLER_MEDSEG.pdf"
