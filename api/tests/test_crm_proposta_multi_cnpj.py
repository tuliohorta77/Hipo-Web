"""
HIPO — 042: oportunidade com vários CNPJs, proposta por CNPJ ou consolidada,
tabela de preço por faixa de vidas.

O que esta suíte segura:

  1. Vincular CNPJ à oportunidade, e as recusas (principal, repetido,
     bloqueado, já em outra negociação aberta).
  2. A conta de cada CNPJ enxerga a oportunidade do grupo (detalhe e
     histórico), e não aceita oportunidade própria enquanto estiver lá.
  3. A proposta com vários CNPJs: soma, desconto contra a tabela, cópia da
     tabela, mensalidade da oportunidade.
  4. O arquivo consolidado e o recortado por CNPJ.
  5. A tabela de preço: leitura para todos, edição só da gestão.
"""
import re
import uuid
from decimal import Decimal
from io import BytesIO

import pytest

from tests.conftest import criar_usuario


def _cnpj(base12: str) -> str:
    def dv(numeros, pesos):
        resto = sum(int(n) * p for n, p in zip(numeros, pesos)) % 11
        return "0" if resto < 2 else str(11 - resto)
    d1 = dv(base12, [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
    d2 = dv(base12 + d1, [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
    return base12 + d1 + d2


MATRIZ = _cnpj("112223330001")
FILIAL_1 = _cnpj("112223330002")
FILIAL_2 = _cnpj("112223330003")
FILIAL_3 = _cnpj("112223330004")
FILIAL_4 = _cnpj("112223330005")
OUTRA = _cnpj("998887770001")

ESCOPO = ["PGR - (NR-01)", "PCMSO - (NR-07)"]


# ── Helpers ──────────────────────────────────────────────────────────

async def nova_conta(client, h, cnpj, razao):
    r = await client.post("/crm/contas", json={"razao_social": razao, "cnpj": cnpj},
                          headers=h)
    assert r.status_code == 201, r.text
    return r.json()


async def nova_oportunidade(client, h, conta_id):
    r = await client.post("/crm/oportunidades", json={"conta_id": conta_id}, headers=h)
    assert r.status_code == 201, r.text
    return r.json()


async def vincular(client, h, opp_id, conta_id):
    return await client.post(f"/crm/oportunidades/{opp_id}/cnpjs",
                             json={"conta_id": conta_id}, headers=h)


async def grupo(client, h, filiais=1):
    """Matriz com oportunidade + N filiais já vinculadas."""
    matriz = await nova_conta(client, h, MATRIZ, "PATIMIRIM PARTICIPACOES LTDA")
    opp = await nova_oportunidade(client, h, matriz["id"])
    contas = [matriz]
    for i, cnpj in enumerate([FILIAL_1, FILIAL_2, FILIAL_3, FILIAL_4][:filiais], start=1):
        c = await nova_conta(client, h, cnpj, f"PATIMIRIM FILIAL {i} LTDA")
        r = await vincular(client, h, opp["id"], c["id"])
        assert r.status_code == 201, r.text
        contas.append(c)
    return opp, contas


def corpo(itens, modalidade="tabela", **troca):
    base = {
        "modalidade": modalidade,
        "itens": itens,
        "escopo": ESCOPO,
        "data_proposta": "2026-09-04",
        "validade": "2026-09-25",
    }
    base.update(troca)
    return base


def _texto_pptx(conteudo: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(BytesIO(conteudo))
    return "\n".join(
        sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame
    )


# ── 1. Vínculo ───────────────────────────────────────────────────────

class TestVinculo:
    async def test_lista_principal_e_adicionais(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=2)
        r = await client.get(f"/crm/oportunidades/{opp['id']}/cnpjs", headers=h)
        assert r.status_code == 200
        lista = r.json()
        assert [c["conta_id"] for c in lista] == [c["id"] for c in contas]
        assert [c["principal"] for c in lista] == [True, False, False]
        assert lista[1]["cnpj_formatado"] == "11.222.333/0002-62"
        assert lista[1]["vinculado_por_nome"]

    async def test_principal_nao_entra_como_adicional(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=0)
        r = await vincular(client, h, opp["id"], contas[0]["id"])
        assert r.status_code == 422
        assert "principal" in r.json()["detail"]

    async def test_repetido_e_recusado(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=1)
        r = await vincular(client, h, opp["id"], contas[1]["id"])
        assert r.status_code == 422
        assert "já está nesta" in r.json()["detail"]

    async def test_cnpj_em_outra_negociacao_aberta_e_409(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=1)
        outra = await nova_conta(client, h, OUTRA, "OUTRO GRUPO LTDA")
        opp2 = await nova_oportunidade(client, h, outra["id"])
        r = await vincular(client, h, opp2["id"], contas[1]["id"])
        assert r.status_code == 409
        d = r.json()["detail"]
        assert d["erro"] == "conta_vinculada_a_oportunidade"
        assert d["numero"] == opp["numero"]
        assert opp["numero"] in d["mensagem"]

    async def test_cnpj_com_oportunidade_propria_aberta_e_409(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, _ = await grupo(client, h, filiais=0)
        outra = await nova_conta(client, h, OUTRA, "OUTRO GRUPO LTDA")
        opp2 = await nova_oportunidade(client, h, outra["id"])
        r = await vincular(client, h, opp["id"], outra["id"])
        assert r.status_code == 409
        assert r.json()["detail"]["numero"] == opp2["numero"]

    async def test_conta_bloqueada_para_prospeccao(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, _ = await grupo(client, h, filiais=0)
        outra = await nova_conta(client, h, OUTRA, "JA E CLIENTE LTDA")
        await db_conn.execute(
            "UPDATE contas SET nao_prospectar = TRUE, nao_prospectar_motivo = 'cliente', "
            "nao_prospectar_em = NOW() "
            "WHERE id = $1", uuid.UUID(outra["id"]),
        )
        r = await vincular(client, h, opp["id"], outra["id"])
        assert r.status_code == 422
        assert r.json()["detail"]["erro"] == "conta_nao_prospectar"

    async def test_desvincular_e_logico(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=1)
        r = await client.delete(f"/crm/oportunidades/{opp['id']}/cnpjs/{contas[1]['id']}",
                                headers=h)
        assert r.status_code == 200
        assert len(r.json()) == 1
        assert await db_conn.fetchval(
            "SELECT count(*) FROM oportunidade_contas WHERE removido_em IS NOT NULL"
        ) == 1
        # E pode voltar.
        assert (await vincular(client, h, opp["id"], contas[1]["id"])).status_code == 201

    async def test_principal_nao_sai(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=0)
        r = await client.delete(f"/crm/oportunidades/{opp['id']}/cnpjs/{contas[0]['id']}",
                                headers=h)
        assert r.status_code == 422

    async def test_resumo_da_oportunidade_conta_os_cnpjs(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, _ = await grupo(client, h, filiais=2)
        r = await client.get(f"/crm/oportunidades/{opp['id']}", headers=h)
        assert r.json()["cnpjs_adicionais"] == 2

    async def test_busca_do_funil_acha_pela_filial(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, _ = await grupo(client, h, filiais=1)
        r = await client.get("/crm/oportunidades", params={"q": FILIAL_1}, headers=h)
        assert [o["id"] for o in r.json()["itens"]] == [opp["id"]]


# ── 2. A conta do CNPJ adicional ─────────────────────────────────────

class TestContaDoCnpj:
    async def test_conta_mostra_a_oportunidade_do_grupo(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=2)
        r = await client.get(f"/crm/contas/{contas[2]['id']}", headers=h)
        [o] = r.json()["oportunidades"]
        assert o["id"] == opp["id"]
        assert o["vinculo"] == "adicional"
        assert o["principal_razao_social"] == "PATIMIRIM PARTICIPACOES LTDA"
        assert o["cnpjs_adicionais"] == 2

    async def test_conta_principal_ve_como_principal(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=1)
        [o] = (await client.get(f"/crm/contas/{contas[0]['id']}", headers=h)).json()["oportunidades"]
        assert o["vinculo"] == "principal"
        assert o["cnpjs_adicionais"] == 1

    async def test_desvinculada_some_da_conta(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=1)
        await client.delete(f"/crm/oportunidades/{opp['id']}/cnpjs/{contas[1]['id']}", headers=h)
        r = await client.get(f"/crm/contas/{contas[1]['id']}", headers=h)
        assert r.json()["oportunidades"] == []

    async def test_historico_da_conta(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=1)
        await client.delete(f"/crm/oportunidades/{opp['id']}/cnpjs/{contas[1]['id']}", headers=h)
        r = await client.get(f"/crm/contas/{contas[1]['id']}/historico", headers=h)
        tipos = [e["tipo"] for e in r.json()]
        assert "cnpj_vinculado" in tipos and "cnpj_desvinculado" in tipos

    async def test_nao_abre_oportunidade_propria(self, db_conn, client, usuario_adm):
        """Decisão do Tulio: bloquear com aviso, apontando a outra."""
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=1)
        r = await client.post("/crm/oportunidades", json={"conta_id": contas[1]["id"]},
                              headers=h)
        assert r.status_code == 409
        assert r.json()["detail"]["oportunidade_id"] == opp["id"]

    async def test_libera_quando_a_do_grupo_fecha(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=1)
        await db_conn.execute(
            "UPDATE oportunidades SET status = 'conquistado', fase = 'finalizado', "
            "fase_desfecho = 'lead' WHERE id = $1", uuid.UUID(opp["id"]),
        )
        r = await client.post("/crm/oportunidades", json={"conta_id": contas[1]["id"]},
                              headers=h)
        assert r.status_code == 201

    async def test_trocar_a_principal_para_um_adicional_dela_mesma(
        self, db_conn, client, usuario_adm
    ):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=1)
        r = await client.patch(f"/crm/oportunidades/{opp['id']}",
                               json={"conta_id": contas[1]["id"]}, headers=h)
        assert r.status_code == 422
        assert "adicional desta oportunidade" in r.json()["detail"]


# ── 3. Proposta com vários CNPJs ─────────────────────────────────────

class TestPropostaVariosCnpjs:
    async def _material(self, client, h):
        """Os cinco CNPJs do material "VARIOS CNPJs": total R$ 1.270,00."""
        opp, contas = await grupo(client, h, filiais=4)
        itens = [
            {"conta_id": contas[0]["id"], "vidas": 4},
            {"conta_id": contas[1]["id"], "vidas": 5},
            {"conta_id": contas[2]["id"], "vidas": 11},
            {"conta_id": contas[3]["id"], "vidas": 23, "mensalidade": "299.00"},
            {"conta_id": contas[4]["id"], "vidas": 27, "mensalidade": "351.00"},
        ]
        r = await client.post(f"/crm/oportunidades/{opp['id']}/propostas",
                              json=corpo(itens), headers=h)
        assert r.status_code == 201, r.text
        return opp, contas, r.json()

    async def test_soma_dos_cnpjs(self, db_conn, client, usuario_adm):
        _, _, p = await self._material(client, usuario_adm["headers"])
        assert p["modalidade"] == "tabela"
        assert Decimal(p["mensalidade"]) == Decimal("1270.00")
        assert p["vidas"] == 70
        assert p["valor_por_vida"] is None
        assert [Decimal(i["mensalidade"]) for i in p["itens"]] == [
            Decimal("180"), Decimal("180"), Decimal("260"), Decimal("299"), Decimal("351"),
        ]

    async def test_desconto_contra_a_tabela(self, db_conn, client, usuario_adm):
        _, _, p = await self._material(client, usuario_adm["headers"])
        quarto = p["itens"][3]
        assert Decimal(quarto["valor_tabela"]) == Decimal("345.00")
        assert Decimal(quarto["desconto_percentual"]) == Decimal("13.3")
        assert p["itens"][0]["desconto_percentual"] is None

    async def test_mensalidade_da_oportunidade_e_a_soma(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, _, _ = await self._material(client, h)
        r = await client.get(f"/crm/oportunidades/{opp['id']}", headers=h)
        assert Decimal(r.json()["valor_mensalidade"]) == Decimal("1270.00")

    async def test_cnpj_fora_da_oportunidade_e_422(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=0)
        solta = await nova_conta(client, h, OUTRA, "SOLTA LTDA")
        r = await client.post(
            f"/crm/oportunidades/{opp['id']}/propostas",
            json=corpo([{"conta_id": contas[0]["id"], "vidas": 4},
                        {"conta_id": solta["id"], "vidas": 4}]),
            headers=h,
        )
        assert r.status_code == 422
        assert "não está vinculado" in r.json()["detail"]

    async def test_por_vida_com_varios_cnpjs(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=1)
        r = await client.post(
            f"/crm/oportunidades/{opp['id']}/propostas",
            json=corpo([{"conta_id": contas[0]["id"], "vidas": 10, "mensalidade": "1"},
                        {"conta_id": contas[1]["id"], "vidas": 5}],
                       modalidade="por_vida", valor_por_vida="20.00",
                       treinamentos="100.00"),
            headers=h,
        )
        assert r.status_code == 201, r.text
        p = r.json()
        assert Decimal(p["mensalidade"]) == Decimal("300.00")   # 15 x 20, o "1" é ignorado
        assert Decimal(p["investimento"]) == Decimal("400.00")
        assert p["tabela_preco"] is None

    async def test_por_vida_sem_valor_e_422(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=0)
        r = await client.post(
            f"/crm/oportunidades/{opp['id']}/propostas",
            json=corpo([{"conta_id": contas[0]["id"], "vidas": 10}], modalidade="por_vida"),
            headers=h,
        )
        assert r.status_code == 422

    async def test_proposta_guarda_a_tabela_do_dia(self, db_conn, client, usuario_adm):
        """Reajuste depois não muda proposta enviada."""
        h = usuario_adm["headers"]
        opp, _, p = await self._material(client, h)
        r = await client.put("/crm/tabela-precos", json={"faixas": [
            {"vidas_ate": 5, "tipo": "fixo", "valor": "999.00"},
            {"vidas_ate": None, "tipo": "por_vida", "valor": "50.00"},
        ]}, headers=h)
        assert r.status_code == 200
        [lida] = (await client.get(f"/crm/oportunidades/{opp['id']}/propostas",
                                   headers=h)).json()
        assert Decimal(lida["tabela_preco"][0]["valor"]) == Decimal("180.00")
        assert Decimal(lida["mensalidade"]) == Decimal("1270.00")

    async def test_padrao_repete_os_cnpjs_da_ultima(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas, _ = await self._material(client, h)
        d = (await client.get(f"/crm/oportunidades/{opp['id']}/proposta-padrao",
                              headers=h)).json()
        assert d["modalidade"] == "tabela"
        assert len(d["cnpjs"]) == 5
        assert [i["vidas"] for i in d["ultimos_itens"]] == [4, 5, 11, 23, 27]
        assert len(d["tabela"]["linhas"]) == 5

    async def test_sem_proposta_o_padrao_e_tabela(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, _ = await grupo(client, h, filiais=0)
        d = (await client.get(f"/crm/oportunidades/{opp['id']}/proposta-padrao",
                              headers=h)).json()
        assert d["modalidade"] == "tabela"
        assert d["ultimos_itens"] == []
        assert d["tabela"]["padrao"] is True


# ── 4. Arquivos ──────────────────────────────────────────────────────

class TestArquivos:
    async def _proposta(self, client, h):
        opp, contas = await grupo(client, h, filiais=2)
        itens = [
            {"conta_id": contas[0]["id"], "vidas": 4},
            {"conta_id": contas[1]["id"], "vidas": 11},
            {"conta_id": contas[2]["id"], "vidas": 23, "mensalidade": "299.00"},
        ]
        p = (await client.post(f"/crm/oportunidades/{opp['id']}/propostas",
                               json=corpo(itens, treinamentos="500.00"), headers=h)).json()
        return opp, p

    async def test_consolidada(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, p = await self._proposta(client, h)
        r = await client.get(f"/crm/propostas/{p['id']}/arquivo", headers=h)
        assert r.status_code == 200, r.text
        texto = _texto_pptx(r.content)
        assert not re.findall(r"\{\{\w+\}\}", texto)
        # 051: a tabela de preços NÃO vai para o cliente; cada CNPJ mostra
        # a faixa em que foi enquadrado.
        assert "TABELA DE PREÇOS" not in texto
        assert "funcionários registrados" not in texto
        assert "(11.222.333/0001-81) - 1 a 5 vidas - Mensalidade R$ 180,00" in texto
        assert "(11.222.333/0002-62) - 11 a 15 vidas - Mensalidade R$ 260,00" in texto
        assert "- 23 vidas - Mensalidade R$ 299,00" in texto   # acima das faixas
        assert "Mensalidade total para os 3 CNPJs -" in texto
        assert "R$ 739,00" in texto                       # 180 + 260 + 299
        assert "será acrescido o valor de R$ 15,00 mensais" in texto
        assert "QTDE. VIDAS: 38" in texto
        assert "Treinamentos - R$ 500,00" in texto
        assert "INVESTIMENTO" not in texto                # slide por vida saiu
        assert "PATIMIRIM PARTICIPACOES LTDA" in texto
        assert opp["numero"] in r.headers["content-disposition"]

    async def test_por_cnpj(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        _, p = await self._proposta(client, h)
        item = p["itens"][1]
        r = await client.get(f"/crm/propostas/{p['id']}/arquivo",
                             params={"item": item["id"]}, headers=h)
        assert r.status_code == 200, r.text
        texto = _texto_pptx(r.content)
        assert not re.findall(r"\{\{\w+\}\}", texto)
        assert "PATIMIRIM FILIAL 1 LTDA" in texto         # cliente = o CNPJ
        assert "QTDE. VIDAS: 11" in texto
        assert "R$ 260,00" in texto
        assert "R$ 739,00" not in texto                   # nada do consolidado
        assert "Treinamentos" not in texto
        assert "Mensalidade R$ 180,00" not in texto       # sem a linha dos outros
        assert "- 11 a 15 vidas - Mensalidade R$ 260,00" in texto   # só a dele
        assert "Mensalidade total" not in texto
        assert item["cnpj"] in r.headers["content-disposition"]

    async def test_item_de_outra_proposta_e_404(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        _, p = await self._proposta(client, h)
        r = await client.get(f"/crm/propostas/{p['id']}/arquivo",
                             params={"item": str(uuid.uuid4())}, headers=h)
        assert r.status_code == 404

    async def test_por_vida_consolidada_mantem_o_quadro(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=1)
        p = (await client.post(
            f"/crm/oportunidades/{opp['id']}/propostas",
            json=corpo([{"conta_id": contas[0]["id"], "vidas": 10},
                        {"conta_id": contas[1]["id"], "vidas": 5}],
                       modalidade="por_vida", valor_por_vida="20.00"),
            headers=h)).json()
        texto = _texto_pptx((await client.get(f"/crm/propostas/{p['id']}/arquivo",
                                              headers=h)).content)
        assert "INVESTIMENTO" in texto
        assert "TABELA DE PREÇOS" not in texto
        assert "R$ 300,00" in texto
        assert "PATIMIRIM FILIAL 1 LTDA (11.222.333/0002-62) - 5 vidas" in texto


# ── 5. Tabela de preço ───────────────────────────────────────────────

class TestTabelaDePreco:
    async def test_banco_vazio_vale_a_padrao(self, db_conn, client, usuario_adm):
        r = await client.get("/crm/tabela-precos", headers=usuario_adm["headers"])
        d = r.json()
        assert d["padrao"] is True
        assert [f["vidas_ate"] for f in d["faixas"]] == [5, 10, 15, 20, None]
        assert d["pode_editar"] is True

    async def test_gestao_salva(self, db_conn, client, usuario_franqueado):
        h = usuario_franqueado["headers"]
        r = await client.put("/crm/tabela-precos", json={"faixas": [
            {"vidas_ate": None, "tipo": "por_vida", "valor": "16.00"},
            {"vidas_ate": 10, "tipo": "fixo", "valor": "230.00"},
        ]}, headers=h)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["padrao"] is False
        assert [f["vidas_ate"] for f in d["faixas"]] == [10, None]
        assert d["atualizado_por_nome"]
        assert d["excedente_padrao"] == "16.00"

    @pytest.mark.parametrize("cargo", ["EV", "EC", "SDR", "EP"])
    async def test_operacional_le_mas_nao_edita(self, db_conn, client, cargo):
        u = await criar_usuario(db_conn, client, cargo, f"{cargo.lower()}-tab@teste.com")
        r = await client.get("/crm/tabela-precos", headers=u["headers"])
        assert r.status_code == 200
        assert r.json()["pode_editar"] is False
        r = await client.put("/crm/tabela-precos", json={"faixas": [
            {"vidas_ate": None, "tipo": "por_vida", "valor": "1.00"},
        ]}, headers=u["headers"])
        assert r.status_code == 403

    async def test_tabela_invalida_e_422(self, db_conn, client, usuario_adm):
        r = await client.put("/crm/tabela-precos", json={"faixas": [
            {"vidas_ate": 5, "tipo": "fixo", "valor": "180.00"},
        ]}, headers=usuario_adm["headers"])
        assert r.status_code == 422
        assert "sem limite" in r.json()["detail"]

    async def test_proposta_nova_usa_a_tabela_salva(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        await client.put("/crm/tabela-precos", json={"faixas": [
            {"vidas_ate": 10, "tipo": "fixo", "valor": "250.00"},
            {"vidas_ate": None, "tipo": "por_vida", "valor": "20.00"},
        ]}, headers=h)
        opp, contas = await grupo(client, h, filiais=0)
        p = (await client.post(f"/crm/oportunidades/{opp['id']}/propostas",
                               json=corpo([{"conta_id": contas[0]["id"], "vidas": 4}]),
                               headers=h)).json()
        assert Decimal(p["mensalidade"]) == Decimal("250.00")


# ── 6. Prospecção enxerga o CNPJ adicional ───────────────────────────

class TestProspeccao:
    async def test_cnpj_adicional_conta_como_em_negociacao(self, db_conn, client, usuario_adm):
        from tests.test_crm_prospeccao import BETA, _semear_base

        h = usuario_adm["headers"]
        await _semear_base(db_conn)
        matriz = await nova_conta(client, h, MATRIZ, "GRUPO LTDA")
        opp = await nova_oportunidade(client, h, matriz["id"])
        beta = await nova_conta(client, h, BETA, "CONSTRUTORA BETA LTDA")
        assert (await vincular(client, h, opp["id"], beta["id"])).status_code == 201

        r = await client.get("/crm/prospeccao",
                             params={"uf": "SP", "cnae": "41", "situacao": "todas"},
                             headers=h)
        assert r.status_code == 200, r.text
        sit = {i["cnpj"]: i["situacao"] for i in r.json()["itens"]}
        assert sit[BETA] == "em_negociacao"


# ── 7. 051: valor por vida excedente e aprovação ─────────────────────

class TestExcedente:
    async def test_sem_informar_vem_da_tabela(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=0)
        p = (await client.post(f"/crm/oportunidades/{opp['id']}/propostas",
                               json=corpo([{"conta_id": contas[0]["id"], "vidas": 4}]),
                               headers=h)).json()
        assert Decimal(p["valor_vida_excedente"]) == Decimal("15.00")

    async def test_ev_escolhe_e_sai_no_arquivo(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=0)
        p = (await client.post(
            f"/crm/oportunidades/{opp['id']}/propostas",
            json=corpo([{"conta_id": contas[0]["id"], "vidas": 18}],
                       valor_vida_excedente="12.50"),
            headers=h)).json()
        assert Decimal(p["valor_vida_excedente"]) == Decimal("12.50")
        texto = _texto_pptx((await client.get(f"/crm/propostas/{p['id']}/arquivo",
                                              headers=h)).content)
        assert "será acrescido o valor de R$ 12,50 mensais" in texto
        assert "- 16 a 20 vidas - Mensalidade R$ 300,00" in texto
        assert "Mensalidade R$ 300,00" in texto              # um CNPJ: rótulo curto

    async def test_padrao_repete_o_da_ultima(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=0)
        await client.post(
            f"/crm/oportunidades/{opp['id']}/propostas",
            json=corpo([{"conta_id": contas[0]["id"], "vidas": 4}], valor_vida_excedente="13"),
            headers=h)
        d = (await client.get(f"/crm/oportunidades/{opp['id']}/proposta-padrao",
                              headers=h)).json()
        assert Decimal(d["valor_vida_excedente"]) == Decimal("13.00")

    async def test_tabela_sem_faixa_por_vida_exige_o_campo(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        await client.put("/crm/tabela-precos", json={"faixas": [
            {"vidas_ate": 10, "tipo": "fixo", "valor": "200.00"},
            {"vidas_ate": None, "tipo": "fixo", "valor": "500.00"},
        ]}, headers=h)
        opp, contas = await grupo(client, h, filiais=0)
        r = await client.post(f"/crm/oportunidades/{opp['id']}/propostas",
                              json=corpo([{"conta_id": contas[0]["id"], "vidas": 4}]),
                              headers=h)
        assert r.status_code == 422
        assert "excedente" in r.json()["detail"]

    async def test_por_vida_ignora_o_campo(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=0)
        p = (await client.post(
            f"/crm/oportunidades/{opp['id']}/propostas",
            json=corpo([{"conta_id": contas[0]["id"], "vidas": 4}], modalidade="por_vida",
                       valor_por_vida="20.00", valor_vida_excedente="99"),
            headers=h)).json()
        assert p["valor_vida_excedente"] is None


class TestAprovacao:
    async def _gerada(self, client, h):
        opp, contas = await grupo(client, h, filiais=0)
        r = await client.post(f"/crm/oportunidades/{opp['id']}/propostas",
                              json=corpo([{"conta_id": contas[0]["id"], "vidas": 4}]),
                              headers=h)
        return opp, r.json()

    async def test_nasce_sem_aprovacao(self, db_conn, client, usuario_adm):
        _, p = await self._gerada(client, usuario_adm["headers"])
        assert p["aprovada_em"] is None
        assert p["aprovada_por_nome"] is None

    async def test_aprovar_grava_quem_e_quando(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, p = await self._gerada(client, h)
        r = await client.post(f"/crm/propostas/{p['id']}/aprovar", headers=h)
        assert r.status_code == 200, r.text
        assert r.json()["aprovada_em"]
        assert r.json()["aprovada_por_nome"]
        [lida] = (await client.get(f"/crm/oportunidades/{opp['id']}/propostas",
                                   headers=h)).json()
        assert lida["aprovada_em"] == r.json()["aprovada_em"]

    async def test_aprovar_de_novo_nao_troca_quem_aprovou(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        _, p = await self._gerada(client, h)
        primeira = (await client.post(f"/crm/propostas/{p['id']}/aprovar", headers=h)).json()
        outro = await criar_usuario(db_conn, client, "EV", "ev-aprova@teste.com")
        segunda = (await client.post(f"/crm/propostas/{p['id']}/aprovar",
                                     headers=outro["headers"])).json()
        assert segunda["aprovada_em"] == primeira["aprovada_em"]
        assert segunda["aprovada_por_nome"] == primeira["aprovada_por_nome"]

    async def test_versao_nova_nasce_sem_aprovacao(self, db_conn, client, usuario_adm):
        h = usuario_adm["headers"]
        opp, p = await self._gerada(client, h)
        await client.post(f"/crm/propostas/{p['id']}/aprovar", headers=h)
        conta_id = p["itens"][0]["conta_id"]
        v2 = (await client.post(f"/crm/oportunidades/{opp['id']}/propostas",
                                json=corpo([{"conta_id": conta_id, "vidas": 6}]),
                                headers=h)).json()
        assert v2["versao"] == 2 and v2["aprovada_em"] is None

    async def test_inexistente_404(self, db_conn, client, usuario_adm):
        r = await client.post(f"/crm/propostas/{uuid.uuid4()}/aprovar",
                              headers=usuario_adm["headers"])
        assert r.status_code == 404


class TestPropriedadesDoArquivo:
    async def test_titulo_do_arquivo_e_do_cliente_certo(self, db_conn, client, usuario_adm):
        """
        051: o modelo carregava o título de outra proposta ("SOLAR DOS
        PAMPAS") e todo PDF saía com o nome de outro cliente na barra.
        """
        from pptx import Presentation

        h = usuario_adm["headers"]
        opp, contas = await grupo(client, h, filiais=0)
        p = (await client.post(f"/crm/oportunidades/{opp['id']}/propostas",
                               json=corpo([{"conta_id": contas[0]["id"], "vidas": 4}]),
                               headers=h)).json()
        r = await client.get(f"/crm/propostas/{p['id']}/arquivo", headers=h)
        props = Presentation(BytesIO(r.content)).core_properties
        assert props.title == "Proposta Comercial - PATIMIRIM PARTICIPACOES LTDA"
        assert "SOLAR" not in (props.title or "")
        assert props.author == p["executivo_nome"]
