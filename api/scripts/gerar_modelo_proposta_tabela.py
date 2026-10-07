"""
Acrescenta ao modelo da proposta o slide da modalidade "tabela por faixa"
(entregas 042 e 051).

    python -m scripts.gerar_modelo_proposta_tabela "PROPOSTA ... VARIOS CNPJs.pptx"

Parte do modelo SEM os slides da tabela (o da 041, ou o que sair do
gerar_modelo_proposta.py numa arte nova) e devolve 7 slides:

    1-4  institucionais (não mudam)
    5    escopo + quadro de investimento   -> modalidade por_vida
    6    escopo + uma linha por CNPJ + mensalidade total -> modalidade tabela
    7    fechamento

O render (services/proposta_render.py) apaga o slide da modalidade que não
foi escolhida. Cada slide variável leva uma etiqueta no NOME de um shape
(`hipo-modalidade:por_vida` / `hipo-modalidade:tabela`).

## 051: a tabela de preços NÃO vai para o cliente

Até a 042 havia um slide "Tabela de preços" (do material). Saiu: a tabela
é a base do vendedor. O cliente vê, em cada linha de CNPJ, a faixa em que
ele foi enquadrado ("16 a 20 vidas - Mensalidade R$ 300,00"), a frase
"Mensalidade total para os N CNPJs - R$ X" e, no rodapé, o valor por vida
excedente que o EV escolheu ({{VALOR_EXCEDENTE}}).

## Por que copiar do material, e não desenhar

O slide 6 vem do material "VARIOS CNPJs" da Controller MedSeg: mesma
moldura, mesmas faixas verdes. O fundo usa exatamente as mesmas imagens do
slide 5 do modelo (conferido por hash), então a cópia não acrescenta mídia
nenhuma — o arquivo não engorda.

## Um run por marcador

Mesma regra do gerar_modelo_proposta.py: cada marcador vive num run único,
senão nenhum replace de string o encontra.
"""
from __future__ import annotations

import copy
import hashlib
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

RAIZ = Path(__file__).resolve().parent.parent
MODELO = RAIZ / "templates" / "proposta_modelo.pptx"

ETIQUETA = "hipo-modalidade:"

# Índices no MODELO atual (0-based).
SLIDE_POR_VIDA = 4
# Índice no material "VARIOS CNPJs" (o slide do escopo com os CNPJs).
ORIGEM_ESCOPO = 5


def _shape(slide, nome):
    for sh in slide.shapes:
        if sh.name == nome:
            return sh
    raise KeyError(f"shape '{nome}' não existe no slide")


def _hash(part) -> str:
    return hashlib.sha1(part.blob).hexdigest()


def _run_unico(paragrafo, texto: str) -> None:
    runs = paragrafo.runs
    if not runs:
        raise ValueError("parágrafo sem runs — nada de onde herdar formatação")
    runs[0].text = texto
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def _texto_unico(shape, texto: str) -> None:
    """Um parágrafo, um run, com a formatação do primeiro run original."""
    tf = shape.text_frame
    _run_unico(tf.paragraphs[0], texto)
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)


def _copiar_slide(destino, origem_slide, layout, por_hash: dict):
    """
    Cria no destino um slide com a árvore de shapes da origem.

    Cada imagem referenciada é religada à parte EQUIVALENTE do destino
    (mesmo conteúdo, por hash). Imagem sem equivalente aborta: copiar
    mídia nova não é o objetivo, e ela engordaria o modelo em silêncio.
    """
    novo = destino.slides.add_slide(layout)
    arvore = novo.shapes._spTree
    for filho in list(arvore):
        if filho.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            arvore.remove(filho)

    origem_arvore = copy.deepcopy(origem_slide.shapes._spTree)

    # A imagem da tabela (Imagem 9) não vem.
    for pic in origem_arvore.findall(qn("p:pic")):
        origem_arvore.remove(pic)

    rels_origem = origem_slide.part.rels
    for el in origem_arvore.iter():
        for atributo in (qn("r:embed"), qn("r:id"), qn("r:link")):
            rid = el.get(atributo)
            if not rid:
                continue
            rel = rels_origem[rid]
            if rel.is_external:
                novo_rid = novo.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                alvo = por_hash.get(_hash(rel.target_part))
                if alvo is None:
                    raise RuntimeError(
                        f"imagem {rel.target_part.partname} do material não existe "
                        "no modelo — o fundo mudou? Confira antes de seguir."
                    )
                novo_rid = novo.part.relate_to(alvo, rel.reltype)
            el.set(atributo, novo_rid)

    for filho in origem_arvore:
        if filho.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        arvore.append(filho)

    # O fundo verde é um gradiente do PRÓPRIO slide (p:bg), fora da árvore
    # de shapes. Sem ele o slide sai branco e o rodapé, que é texto branco,
    # some.
    fundo = origem_slide._element.cSld.find(qn("p:bg"))
    if fundo is not None:
        csld = novo._element.cSld
        antigo = csld.find(qn("p:bg"))
        if antigo is not None:
            csld.remove(antigo)
        csld.insert(0, copy.deepcopy(fundo))
    return novo


def _mover(prs, slide, para_indice: int) -> None:
    lista = prs.slides._sldIdLst
    alvo = None
    for sld in lista:
        if prs.part.related_part(sld.get(qn("r:id"))) is slide.part:
            alvo = sld
            break
    lista.remove(alvo)
    lista.insert(para_indice, alvo)


def _proximo_id(slide) -> int:
    return max(int(el.get("id")) for el in slide._element.iter(qn("p:cNvPr"))) + 1


def gerar(material: str, modelo: Path = MODELO) -> Path:
    prs = Presentation(str(modelo))
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.name.startswith(ETIQUETA):
                raise SystemExit(
                    "o modelo já tem slides etiquetados — parta do modelo sem "
                    "eles (o da 041 ou o do gerar_modelo_proposta.py)."
                )

    origem = Presentation(material)
    s5 = prs.slides[SLIDE_POR_VIDA]
    layout = s5.slide_layout

    # Toda imagem que o modelo já tem, por conteúdo.
    por_hash = {}
    for slide in prs.slides:
        for rel in slide.part.rels.values():
            if not rel.is_external and rel.reltype.endswith("/image"):
                por_hash[_hash(rel.target_part)] = rel.target_part

    molde_escopo = _shape(s5, "TextBox 8").text_frame._txBody

    escopo = _copiar_slide(prs, origem.slides[ORIGEM_ESCOPO], layout, por_hash)
    caixa = _shape(escopo, "TextBox 8")
    caixa._element.replace(caixa._element.txBody, copy.deepcopy(molde_escopo))
    caixa.name = f"{ETIQUETA}tabela"
    _texto_unico(_shape(escopo, "TextBox 16"), "QTDE. VIDAS: {{VIDAS}}")

    # Mensalidade: uma caixa só, da largura do quadro, centralizada —
    # "Mensalidade total para os 5 CNPJs - R$ 4.500,00". Rótulo e valor
    # continuam em runs separados, cada um com a formatação do material.
    rotulo = _shape(escopo, "TextBox 25")
    valor = _shape(escopo, "TextBox 24")
    run_valor = copy.deepcopy(valor.text_frame.paragraphs[0].runs[0]._r)
    _texto_unico(rotulo, "{{ROTULO_MENSALIDADE}} ")
    paragrafo = rotulo.text_frame.paragraphs[0]
    paragrafo._p.append(run_valor)
    paragrafo.runs[1].text = "{{MENSALIDADE}}"
    from pptx.enum.text import PP_ALIGN
    paragrafo.alignment = PP_ALIGN.CENTER
    rotulo.left, rotulo.width = caixa.left, caixa.width
    rotulo.text_frame.word_wrap = False
    rotulo.name = "hipo-mensalidade"
    for vazio in ("TextBox 18", "TextBox 19", "TextBox 24"):
        sh = _shape(escopo, vazio)
        sh._element.getparent().remove(sh._element)

    # Rodapé: o mesmo do slide por vida, com o excedente que o EV escolhe.
    rodape = copy.deepcopy(_shape(s5, "TextBox 15")._element)
    rodape.find(f"{qn('p:nvSpPr')}/{qn('p:cNvPr')}").set("id", str(_proximo_id(escopo)))
    escopo.shapes._spTree.append(rodape)
    rodape_shape = escopo.shapes[-1]
    run = rodape_shape.text_frame.paragraphs[0].runs[0]
    if "{{VALOR_VIDA}}" not in run.text:
        raise RuntimeError("o rodapé do slide por vida mudou — confira o modelo.")
    run.text = run.text.replace("{{VALOR_VIDA}}", "{{VALOR_EXCEDENTE}}")

    # Etiqueta do slide por vida.
    _shape(s5, "TextBox 8").name = f"{ETIQUETA}por_vida"

    # Ordem: 1-4, 5 por vida, 6 tabela, 7 fechamento.
    _mover(prs, escopo, SLIDE_POR_VIDA + 1)

    prs.save(str(modelo))
    return modelo


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)
    print(f"Modelo atualizado: {gerar(sys.argv[1])}")
