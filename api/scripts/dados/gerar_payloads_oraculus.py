"""
HIPO — Transformação XLSX → JSON das duas cargas de 09/09/2026.

Este arquivo NÃO fala com o banco. Ele existe para que a normalização seja
auditável e reprodutível: lê a planilha bruta e escreve os dois payloads
(mais os CSVs de conferência) que os importadores consomem.

Mesma decisão da carga do CRM Omie: o que roda contra produção só lê JSON e
grava — sem parsing, sem heurística, sem openpyxl na EC2.

USO
    cd api
    python -m scripts.dados.gerar_payloads_oraculus \
        --xlsx /caminho/Planilha_Geral_Controller_Oraculus_consolidada.xlsx

REGRAS DE NORMALIZAÇÃO (as que custaram análise)

  CNPJ
    Só dígitos. Precisa ter 14 e passar no dígito verificador. Linha que
    falha vai para descartados.csv com o motivo — nenhuma vira placeholder.
    A carga do CRM Omie usou placeholder sequencial porque eram 17 registros
    ATIVOS com histórico comercial; aqui são nomes de uma lista fria, e um
    CNPJ inventado numa base de prospecção só serve para alguém ligar para
    a empresa errada.

  CONTATOS — o pareamento posicional
    As colunas "Contato" e "E-mail" trazem várias pessoas separadas por " / ",
    e as duas listas estão ALINHADAS POSIÇÃO A POSIÇÃO: 419 das 422 linhas
    com contato têm exatamente a mesma quantidade de fatias nas duas colunas.
    Então o nome da fatia i pertence ao e-mail da fatia i.

    Isso é o que permite honrar "um contato por e-mail" sem inventar par: a
    tentação era casar por ordem de e-mail e sobrar nome, e aí "Alessandra"
    ganharia o e-mail do "Andre".

    Fatia de nome que é ela mesma um endereço (89 casos) vira contato com o
    e-mail no lugar do nome. Não se deriva nome do local part: "adm" e
    "financeiro" não são pessoas, e um nome inventado é pior que um e-mail
    honesto na tela.

    As 3 linhas desalinhadas: 2 sem nome (o e-mail vira o nome) e 1 com dois
    nomes para um e-mail só (o primeiro fica com o e-mail, o segundo entra
    sem e-mail). Ficam listadas em contatos_desalinhados.csv.

  CONTATO REPETIDO
    Dedup por e-mail em minúsculas, e SÓ por e-mail. Nunca por nome: a carga
    do CRM Omie já registrou que casar por nome funde homônimos de empresas
    diferentes. Contato sem e-mail é sempre um contato novo.

  Nº DE FUNCIONÁRIOS
    Três colunas na planilha (Pessoal, Página21, Folha Edivaldo). Vale a
    primeira que for um número inteiro, nessa ordem. Valor não numérico
    ("ALECIO") é ignorado.

  O QUE NÃO É IMPORTADO
    Regime tributário, I.M., I.E., certificado, divisão contábil, sindicato:
    são dados da operação CONTÁBIL da Oraculus, não do nosso funil. O regime
    tributário vai para observacoes como texto, porque é o único com chance
    de virar critério comercial — os outros nem isso.

    UF fica NULL nas 923 linhas em que a planilha não trouxe, mesmo com a
    cidade preenchida. Deduzir "GUARULHOS → SP" é enriquecimento, e a base
    passaria a afirmar como dado o que foi palpite nosso.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

import openpyxl

AQUI = Path(__file__).resolve().parent

# CNPJ da aba MedSeg que NAO deve ser bloqueado nesta rodada, com o motivo.
#
# Fica aqui, e nao removido da planilha na mao, porque decisao de negocio que
# nao esta escrita em lugar nenhum vira mistério na proxima carga: alguem
# regeraria o JSON a partir da planilha original e o CNPJ voltaria calado.
# Sai tambem em medseg_fora.csv, para nao se perder.
MEDSEG_FORA = {
    "09006006000140": (
        "AERO PALETES tem oportunidade em aberto; Tulio decidiu em 09/09/2026 "
        "encerrar o negocio antes de bloquear. Bloquear a mao pela tela depois."
    ),
}

# ── CNPJ (cópia local das puras de services/cnpj.py, para o gerador rodar
#    sozinho, fora do pacote da API) ──────────────────────────────────────
_SO_DIGITOS = re.compile(r"\D")
_REPETIDOS = {str(d) * 14 for d in range(10)}
_PESOS_1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
_PESOS_2 = [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]


def so_digitos(v) -> str:
    return _SO_DIGITOS.sub("", str(v)) if v is not None else ""


def _dv(base: str, pesos: list[int]) -> str:
    resto = sum(int(d) * p for d, p in zip(base, pesos)) % 11
    return "0" if resto < 2 else str(11 - resto)


def cnpj_valido(num: str) -> bool:
    if len(num) != 14 or num in _REPETIDOS:
        return False
    return num[12] == _dv(num[:12], _PESOS_1) and num[13] == _dv(num[:13], _PESOS_2)


_ESPACOS = re.compile(r"\s+")


def limpar(v) -> str:
    if v is None:
        return ""
    return _ESPACOS.sub(" ", str(v)).strip()


def sem_acento_maiusculo(v: str) -> str:
    s = unicodedata.normalize("NFKD", v)
    return "".join(c for c in s if not unicodedata.combining(c)).upper().strip()


def fatias(v) -> list[str]:
    """Quebra o campo pelo separador da planilha (' / '), mantendo a posição."""
    s = limpar(v)
    if not s:
        return []
    return [p.strip() for p in re.split(r"\s*/\s*", s)]


EH_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def inteiro(v):
    s = limpar(v)
    if not s:
        return None
    s = s.replace(".", "").replace(",", "")
    return int(s) if s.isdigit() else None


def telefone(v) -> str | None:
    """A coluna guarda '11 94882-7127'. A do banco é VARCHAR(20)."""
    s = limpar(v)
    return s[:20] if s else None


# ── Oraculus ─────────────────────────────────────────────────────────────

def montar_oraculus(ws, hoje: str) -> tuple[dict, list, list]:
    cab = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {h: n for n, h in enumerate(cab)}
    linhas = [r for r in ws.iter_rows(min_row=2, values_only=True)
              if any(c not in (None, "") for c in r)]

    registros: list[dict] = []
    descartados: list[dict] = []
    desalinhados: list[dict] = []
    vistos: set[str] = set()

    for n, r in enumerate(linhas, start=2):
        bruto = r[col["CNPJ / CPF"]]
        razao = limpar(r[col["Empresa / Razão Social"]])
        num = so_digitos(bruto)

        if not cnpj_valido(num):
            descartados.append({
                "linha": n,
                "documento": limpar(bruto),
                "tipo_doc": limpar(r[col["Tipo Doc"]]),
                "razao_social": razao,
                "motivo": ("CPF, nao cabe em contas.cnpj" if len(num) == 11
                           else "CNPJ invalido (digito verificador)"),
            })
            continue
        if not razao:
            descartados.append({"linha": n, "documento": limpar(bruto),
                                "tipo_doc": "", "razao_social": "",
                                "motivo": "sem razao social"})
            continue
        if num in vistos:
            descartados.append({"linha": n, "documento": limpar(bruto),
                                "tipo_doc": limpar(r[col["Tipo Doc"]]),
                                "razao_social": razao,
                                "motivo": "CNPJ repetido na propria planilha"})
            continue
        vistos.add(num)

        # ── contatos, pareados posição a posição ──────────────────────
        nomes = fatias(r[col["Contato"]])
        emails = [e.lower() for e in fatias(r[col["E-mail"]]) if EH_EMAIL.match(e.lower())]
        if len(nomes) != len(emails):
            desalinhados.append({
                "linha": n, "cnpj": num, "razao_social": razao,
                "contato_bruto": limpar(r[col["Contato"]]),
                "email_bruto": limpar(r[col["E-mail"]]),
                "qtd_nomes": len(nomes), "qtd_emails": len(emails),
            })

        fone = telefone(r[col["Telefone"]])
        contatos = []
        for i in range(max(len(nomes), len(emails))):
            nome_i = nomes[i] if i < len(nomes) else ""
            email_i = emails[i] if i < len(emails) else None
            # Fatia de nome que é um endereço: o e-mail vira o nome exibido.
            if EH_EMAIL.match(nome_i.lower()):
                nome_i = ""
            nome_final = nome_i or (email_i or "")
            if not nome_final:
                continue
            contatos.append({
                "nome": nome_final[:150],
                "email": email_i,
                # O telefone da planilha é da EMPRESA, não da pessoa. Vai só
                # para contas.telefone; repetir na pessoa afirmaria que
                # aquele número é dela.
                "telefone": None,
                "principal": len(contatos) == 0,
            })

        obs = []
        regime = limpar(r[col["Regime Tributário"]])
        if regime:
            obs.append(f"Regime tributario: {regime}")
        codigo = limpar(r[col["Código"]])
        if codigo:
            obs.append(f"Codigo Oraculus: {codigo}")

        registros.append({
            "chave": num,
            "linha": n,
            "conta": {
                "cnpj": num,
                "razao_social": razao[:200],
                "cidade": (limpar(r[col["Cidade"]]) or None),
                "uf": (limpar(r[col["UF"]]).upper()[:2] or None),
                "num_funcionarios": (inteiro(r[col["Nº Func. (Pessoal)"]])
                                     or inteiro(r[col["Nº Func. (Página21)"]])
                                     or inteiro(r[col["Nº Func. (Folha Edivaldo)"]])),
                "telefone": fone,
                "observacoes": " | ".join(obs) or None,
            },
            "contatos": contatos,
        })

    payload = {
        "gerado_em": hoje,
        "fonte": "Planilha_Geral_Controller_Oraculus_consolidada.xlsx — aba 'Clientes Oraculus'",
        "origem": {"slug": "carteira-oraculus", "nome": "Carteira Oraculus"},
        "oportunidade_padrao": {"fase": "suspect", "temperatura": 0},
        "registros": registros,
    }
    return payload, descartados, desalinhados


# ── MedSeg ───────────────────────────────────────────────────────────────

def montar_medseg(ws, hoje: str) -> tuple[dict, list]:
    linhas = [r for r in ws.iter_rows(min_row=2, values_only=True)
              if any(c not in (None, "") for c in r)]

    registros: list[dict] = []
    descartados: list[dict] = []
    fora: list[dict] = []
    vistos: dict[str, str] = {}

    for n, r in enumerate(linhas, start=2):
        razao = limpar(r[0])
        num = so_digitos(r[1])
        if num in MEDSEG_FORA:
            fora.append({"linha": n, "cnpj": num, "razao_social": razao,
                         "motivo": MEDSEG_FORA[num]})
            continue
        if not cnpj_valido(num):
            descartados.append({"linha": n, "razao_social": razao,
                                "documento": limpar(r[1]),
                                "motivo": "CNPJ invalido (digito verificador)"})
            continue
        if num in vistos:
            descartados.append({
                "linha": n, "razao_social": razao, "documento": limpar(r[1]),
                "motivo": f"CNPJ repetido na planilha (ja veio como '{vistos[num]}')",
            })
            continue
        vistos[num] = razao
        registros.append({"cnpj": num, "razao_social": razao[:200]})

    payload = {
        "gerado_em": hoje,
        "fonte": "Planilha_Geral_Controller_Oraculus_consolidada.xlsx — aba 'Clientes MEDSEG'",
        "motivo": f"Cliente da Controller MedSeg (carteira consolidada em {hoje})",
        "registros": registros,
    }
    return payload, descartados, fora


def escrever_csv(caminho: Path, linhas: list[dict]) -> None:
    if not linhas:
        caminho.write_text("", encoding="utf-8")
        return
    with caminho.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(linhas[0].keys()))
        w.writeheader()
        w.writerows(linhas)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--saida", default=str(AQUI))
    ap.add_argument("--data", default=date.today().isoformat())
    args = ap.parse_args()

    saida = Path(args.saida)
    saida.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)

    orac, orac_desc, desalin = montar_oraculus(wb["Clientes Oraculus"], args.data)
    med, med_desc, med_fora = montar_medseg(wb["Clientes MEDSEG"], args.data)

    (saida / f"oraculus_{args.data}.json").write_text(
        json.dumps(orac, ensure_ascii=False, indent=1), encoding="utf-8")
    (saida / f"medseg_nao_prospectar_{args.data}.json").write_text(
        json.dumps(med, ensure_ascii=False, indent=1), encoding="utf-8")

    escrever_csv(saida / "oraculus_descartados.csv", orac_desc)
    escrever_csv(saida / "oraculus_contatos_desalinhados.csv", desalin)
    escrever_csv(saida / "medseg_descartados.csv", med_desc)
    escrever_csv(saida / "medseg_fora.csv", med_fora)
    escrever_csv(saida / "conferencia_oraculus.csv", [
        {
            "cnpj": r["conta"]["cnpj"],
            "razao_social": r["conta"]["razao_social"],
            "cidade": r["conta"]["cidade"] or "",
            "uf": r["conta"]["uf"] or "",
            "num_funcionarios": r["conta"]["num_funcionarios"] or "",
            "telefone": r["conta"]["telefone"] or "",
            "qtd_contatos": len(r["contatos"]),
            "contato_principal": (r["contatos"][0]["nome"] if r["contatos"] else ""),
            "email_principal": (r["contatos"][0]["email"] or "" if r["contatos"] else ""),
        }
        for r in orac["registros"]
    ])
    escrever_csv(saida / "conferencia_medseg.csv", med["registros"])

    emails = {c["email"] for r in orac["registros"] for c in r["contatos"] if c["email"]}
    sem_email = sum(1 for r in orac["registros"] for c in r["contatos"] if not c["email"])
    print(f"Oraculus  : {len(orac['registros'])} contas, "
          f"{sum(len(r['contatos']) for r in orac['registros'])} contatos "
          f"({len(emails)} e-mails distintos, {sem_email} sem e-mail), "
          f"{len(orac_desc)} descartadas, {len(desalin)} linhas desalinhadas")
    print(f"MedSeg    : {len(med['registros'])} contas a bloquear, "
          f"{len(med_desc)} descartadas, {len(med_fora)} deixadas de fora")
    for f in med_fora:
        print(f"            fora: {f['cnpj']} {f['razao_social']}")
    cruz = {r["cnpj"] for r in med["registros"]} & {r["chave"] for r in orac["registros"]}
    print(f"Intersecao: {len(cruz)} CNPJ nas duas abas")
    return 0


if __name__ == "__main__":
    sys.exit(main())
