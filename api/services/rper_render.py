"""
HIPO — RPeR: o desenho do .pptx.

O PPT que a operação já usava (RPeR_<MÊS>_2026_CONTROLLER_MEDSEG) foi a
referência de tudo aqui: fundo, logo, as ondas laranja e verde, a
tipografia Poppins, o verde 0E5200 dos cabeçalhos de tabela, a coluna de
META em laranja-claro, os cartões brancos com sombra. As imagens são as
MESMAS do arquivo original, em api/templates/rper/.

POR QUE DESENHAR EM CÓDIGO, E NÃO PREENCHER UM MODELO COM MARCADORES
(como a proposta comercial faz)

A proposta tem campos fixos. O RPeR não: o número de pessoas em cada
squad muda (e com ele as colunas da tabela por pessoa, os cartões de
comentário e as linhas da tabela de metas), o top de negociações pode
ter de zero a dez linhas, e o gráfico tem uma barra por executivo. Um
modelo com marcadores precisaria de clonagem de linha, de COLUNA e de
slide — o python-pptx não clona slide, e clonar coluna de tabela é
cirurgia de XML. Desenhar direto é menos código e nenhum marcador para
alguém apagar sem querer ao abrir o modelo no PowerPoint.

python-pptx é importado dentro das funções, pelo mesmo motivo da
proposta (services/proposta_render.py): import no topo derruba a API
inteira se a biblioteca faltar no servidor.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from services import rper as regras
from services.proposta_render import BibliotecaIndisponivel
from services.tarefa import FUSO_OPERACAO

PASTA_ASSETS = Path(__file__).resolve().parent.parent / "templates" / "rper"

# Paleta do PPT original.
VERDE_ESCURO = "0E5200"
VERDE = "178900"
LARANJA = "F86000"
LARANJA_CLARO = "F69633"
META_FUNDO = "FFF4E8"
ZEBRA = "F4F6F2"
BORDA = "E4E4E4"
TEXTO = "1A1A1A"
TEXTO_2 = "5A5A5A"
CINZA_SUB = "9A9A9A"
BEGE = "E5E2D9"
BRANCO = "FFFFFF"
VERMELHO = "C62828"
VERDE_CLARO_CAPA = "CFE3C8"
VERDE_CLARO_SECAO = "D8E8D0"
CINZA_GRAFICO = "A9BCA4"

FONTE = "Poppins"

# Cor da célula de atingimento pela carinha do Monitor — a mesma régua.
COR_ATINGIMENTO = {
    "muito_feliz": VERDE, "feliz": VERDE,
    "neutro": LARANJA, "triste": VERMELHO, "bravo": VERMELHO,
}

NOME_POR_PESSOA = {
    "EC": "INDICADORES POR EXECUTIVO DE CONTAS",
    "SDR": "INDICADORES POR SDR",
    "EV": "INDICADORES POR EXECUTIVO DE VENDAS",
}

# Os quatro quadros grandes do slide de resultados, por squad.
KPIS = {
    "SDR": ("tarefas", "contas", "taxa_execucao", "pipeline_gerado"),
    "EV": ("taxa_execucao", "propostas", "em_negociacao", "oportunidades"),
    "EC": ("contas_gestao", "reunioes_carteira", "leads", "mrr"),
}
ROTULO_KPI = {
    ("SDR", "tarefas"): "TAREFAS DE\nPROSPECÇÃO",
    ("SDR", "contas"): "CONTAS\nPROSPECTADAS",
    ("SDR", "taxa_execucao"): "TAXA DE\nEXECUÇÃO",
    ("SDR", "pipeline_gerado"): "TICKET GERADO\nNA PROSPECÇÃO",
    ("EV", "taxa_execucao"): "TAXA DE\nEXECUÇÃO",
    ("EV", "propostas"): "PROPOSTAS\nENVIADAS",
    ("EV", "em_negociacao"): "EM\nNEGOCIAÇÃO",
    ("EV", "oportunidades"): "OPORTUNIDADES\nTRABALHADAS",
    ("EC", "contas_gestao"): "CONTAS SOB\nGESTÃO",
    ("EC", "reunioes_carteira"): "REUNIÕES DE\nCARTEIRA",
    ("EC", "leads"): "LEADS\nINDICADOS",
    ("EC", "mrr"): "MRR FECHADO\nCOM EC",
}


def _pptx():
    try:
        import pptx  # noqa: F401
        from pptx import Presentation
    except ImportError as erro:
        raise BibliotecaIndisponivel(
            "python-pptx não está instalado no servidor, então o RPeR não pode "
            "ser gerado. Instale com: "
            "sudo -iu hipo python3 -m pip install --user python-pptx==1.0.2"
        ) from erro
    return Presentation


class _Desenho:
    """Os primitivos visuais do RPeR. Uma instância por arquivo."""

    def __init__(self):
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.oxml.ns import qn

        self.In, self.Pt, self.RGB = Inches, Pt, RGBColor
        self.MSO_SHAPE, self.ALIGN, self.ANCHOR, self.qn = MSO_SHAPE, PP_ALIGN, MSO_ANCHOR, qn

        Presentation = _pptx()
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self._vazio = self.prs.slide_layouts[6]

    # ── básicos ──────────────────────────────────────────────────────

    def slide(self):
        return self.prs.slides.add_slide(self._vazio)

    def imagem(self, slide, nome, x, y, w, h):
        caminho = PASTA_ASSETS / nome
        return slide.shapes.add_picture(
            str(caminho), self.In(x), self.In(y), self.In(w), self.In(h)
        )

    def _cor(self, hexa):
        return self.RGB.from_string(hexa)

    def _run(self, run, texto, *, tam, cor, negrito=False, espaco=None, italico=False):
        run.text = texto
        f = run.font
        f.name = FONTE
        f.size = self.Pt(tam)
        f.bold = negrito
        f.italic = italico
        f.color.rgb = self._cor(cor)
        if espaco is not None:
            run._r.get_or_add_rPr().set("spc", str(espaco))

    def texto(self, slide, x, y, w, h, partes, *, tam=11, cor=TEXTO, negrito=False,
              alinhar="l", ancora="t", espaco=None, entre=0):
        """
        Caixa de texto sem margem interna (como no original). `partes` é
        uma string, uma lista de parágrafos (strings) ou uma lista de
        parágrafos em que cada parágrafo é uma lista de runs
        (texto, {opções}).
        """
        caixa = slide.shapes.add_textbox(self.In(x), self.In(y), self.In(w), self.In(h))
        tf = caixa.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"t": self.ANCHOR.TOP, "ctr": self.ANCHOR.MIDDLE,
                              "b": self.ANCHOR.BOTTOM}[ancora]
        paragrafos = [partes] if isinstance(partes, str) else partes
        for i, par in enumerate(paragrafos):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": self.ALIGN.LEFT, "ctr": self.ALIGN.CENTER,
                           "r": self.ALIGN.RIGHT}[alinhar]
            if entre and i > 0:
                p.space_before = self.Pt(entre)
            runs = [(par, {})] if isinstance(par, str) else par
            for txt, op in runs:
                self._run(
                    p.add_run(), txt,
                    tam=op.get("tam", tam), cor=op.get("cor", cor),
                    negrito=op.get("negrito", negrito),
                    espaco=op.get("espaco", espaco),
                    italico=op.get("italico", False),
                )
        return caixa

    def marcadores(self, slide, x, y, w, h, itens, *, tam=11, cor=TEXTO):
        """Lista com marcador real (buChar), não '•' digitado."""
        caixa = self.texto(slide, x, y, w, h, list(itens) or ["—"], tam=tam, cor=cor, entre=4)
        for p in caixa.text_frame.paragraphs:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(self.In(0.22))))
            pPr.set("indent", str(-int(self.In(0.18))))
            for filho in list(pPr):
                if filho.tag in (self.qn("a:buNone"), self.qn("a:buChar"), self.qn("a:buClr")):
                    pPr.remove(filho)
            buClr = pPr.makeelement(self.qn("a:buClr"), {})
            srgb = buClr.makeelement(self.qn("a:srgbClr"), {"val": LARANJA})
            buClr.append(srgb)
            pPr.append(buClr)
            pPr.append(pPr.makeelement(self.qn("a:buChar"), {"char": "•"}))
        return caixa

    def cartao(self, slide, x, y, w, h, *, fundo=BRANCO, raio=0.06, sombra=True):
        forma = slide.shapes.add_shape(
            self.MSO_SHAPE.ROUNDED_RECTANGLE, self.In(x), self.In(y), self.In(w), self.In(h)
        )
        forma.adjustments[0] = raio
        forma.fill.solid()
        forma.fill.fore_color.rgb = self._cor(fundo)
        forma.line.color.rgb = self._cor(fundo)
        forma.line.width = self.Pt(0.5)
        forma.text_frame.text = ""
        spPr = forma._element.spPr
        if sombra:
            efeitos = spPr.makeelement(self.qn("a:effectLst"), {})
            sh = efeitos.makeelement(self.qn("a:outerShdw"), {
                "blurRad": "152400", "dist": "38100", "dir": "5400000",
                "algn": "bl", "rotWithShape": "0",
            })
            cor = sh.makeelement(self.qn("a:srgbClr"), {"val": "000000"})
            cor.append(cor.makeelement(self.qn("a:alpha"), {"val": "35000"}))
            sh.append(cor)
            efeitos.append(sh)
            spPr.append(efeitos)
        else:
            # Sem isto o tema aplica a sombra padrão da forma.
            spPr.append(spPr.makeelement(self.qn("a:effectLst"), {}))
        return forma

    # ── moldura dos slides ───────────────────────────────────────────

    def fundo_conteudo(self, slide, branco, laranja, subtitulo):
        self.imagem(slide, "bg_conteudo.png", 0, 0, 13.333, 7.5)
        self.imagem(slide, "logo.png", 10.81, 0.24, 2.10, 0.48)
        self.texto(slide, 0.55, 0.30, 10.0, 0.62, [[
            (branco, {"cor": BRANCO}), (laranja, {"cor": LARANJA_CLARO}),
        ]], tam=26, negrito=True, ancora="ctr")
        self.texto(slide, 0.55, 0.90, 10.0, 0.34, subtitulo, tam=12, cor=CINZA_SUB, ancora="ctr")

    def rodape(self, slide, texto):
        self.texto(slide, 0.55, 6.86, 12.2, 0.40, texto, tam=9, cor=CINZA_SUB, ancora="t")

    def rotulo_cartao(self, slide, x, y, w, texto, cor=LARANJA):
        self.texto(slide, x, y, w, 0.30, texto, tam=11, negrito=True, cor=cor, espaco=100)

    def kpi(self, slide, x, y, w, h, valor, rotulo, *, cor=VERDE_ESCURO):
        self.cartao(slide, x, y, w, h, raio=0.11)
        tam = 34 if len(valor) <= 5 else (26 if len(valor) <= 9 else 20)
        self.texto(slide, x, y + 0.16, w, h * 0.46, valor, tam=tam, negrito=True,
                   cor=cor, alinhar="ctr", ancora="ctr")
        self.texto(slide, x + 0.1, y + h * 0.56, w - 0.2, h * 0.40,
                   rotulo.split("\n"), tam=9.5, negrito=True, cor=TEXTO_2,
                   alinhar="ctr", espaco=50)

    # ── tabela ───────────────────────────────────────────────────────

    def _bordas(self, celula):
        tcPr = celula._tc.get_or_add_tcPr()
        # O esquema de tcPr exige lnL, lnR, lnT, lnB ANTES do preenchimento.
        # Inserir no índice 0 em ordem reversa deixa as quatro na frente, na
        # ordem certa — o PowerPoint recusa o arquivo com a ordem trocada.
        for lado in ("a:lnB", "a:lnT", "a:lnR", "a:lnL"):
            ln = tcPr.makeelement(self.qn(lado), {"w": "6350", "cap": "flat", "cmpd": "sng"})
            fill = ln.makeelement(self.qn("a:solidFill"), {})
            fill.append(fill.makeelement(self.qn("a:srgbClr"), {"val": BORDA}))
            ln.append(fill)
            tcPr.insert(0, ln)

    def tabela(self, slide, x, y, larguras, linhas, *, altura_linha=0.36,
               estilos=None, cabecalho=True, zebra=False):
        """
        `linhas`: lista de listas de células. Cada célula é uma string ou
        uma lista de parágrafos [(texto, {tam, cor, negrito})].
        `estilos(i, j)` devolve {fundo, alinhar, tam, cor, negrito} para
        sobrepor o padrão da célula (i conta o cabeçalho como 0).
        """
        n_lin, n_col = len(linhas), len(larguras)
        forma = slide.shapes.add_table(
            n_lin, n_col, self.In(x), self.In(y),
            self.In(sum(larguras)), self.In(altura_linha * n_lin),
        )
        tbl = forma.table
        tblPr = tbl._tbl.tblPr
        for attr in ("firstRow", "bandRow"):
            tblPr.set(attr, "0")
        for j, w in enumerate(larguras):
            tbl.columns[j].width = self.In(w)
        for i in range(n_lin):
            tbl.rows[i].height = self.In(altura_linha)
            for j in range(n_col):
                cel = tbl.cell(i, j)
                eh_cab = cabecalho and i == 0
                est = {
                    "fundo": VERDE_ESCURO if eh_cab else (ZEBRA if zebra and i % 2 == 0 else BRANCO),
                    "alinhar": "l" if j == 0 else "ctr",
                    "tam": 10.5 if eh_cab else 11,
                    "cor": BRANCO if eh_cab else TEXTO,
                    "negrito": True,
                }
                if estilos and not eh_cab:
                    est.update(estilos(i, j) or {})
                cel.margin_left = cel.margin_right = self.In(0.08)
                cel.margin_top = cel.margin_bottom = self.In(0.02)
                cel.vertical_anchor = self.ANCHOR.MIDDLE
                conteudo = linhas[i][j]
                paragrafos = ([[(conteudo, {})]] if isinstance(conteudo, str)
                              else conteudo)
                tf = cel.text_frame
                tf.word_wrap = True
                for k, par in enumerate(paragrafos):
                    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                    p.alignment = {"l": self.ALIGN.LEFT, "ctr": self.ALIGN.CENTER}[est["alinhar"]]
                    for txt, op in par:
                        self._run(p.add_run(), txt, tam=op.get("tam", est["tam"]),
                                  cor=op.get("cor", est["cor"]),
                                  negrito=op.get("negrito", est["negrito"]))
                cel.fill.solid()
                cel.fill.fore_color.rgb = self._cor(est["fundo"])
                self._bordas(cel)
        return forma


# ── Slides ───────────────────────────────────────────────────────────


def _mes_maiusculo(ano, mes):
    return regras.rotulo_mes(ano, mes, com_ano=False).upper()


def _capa(d: _Desenho, r: dict):
    s = d.slide()
    d.imagem(s, "bg_capa.png", 0, 0, 13.333, 7.5)
    d.imagem(s, "swoosh_laranja.png", -2.60, 4.90, 13.00, 2.14)
    d.imagem(s, "swoosh_verde.png", 1.40, 5.50, 13.50, 2.22)
    d.imagem(s, "logo.png", 0.85, 0.85, 4.60, 1.06)
    pilula = d.cartao(s, 0.90, 2.05, 3.20, 0.46, fundo=LARANJA, raio=0.5, sombra=False)
    pilula.name = "Pilula"
    d.texto(s, 0.90, 2.05, 3.20, 0.46,
            f"{_mes_maiusculo(r['ano_novo'], r['mes_novo'])} | {r['ano_novo']}",
            tam=12.5, negrito=True, cor=BRANCO, alinhar="ctr", ancora="ctr")
    for i, (linha, cor) in enumerate((("REUNIÃO DE", BRANCO), ("PLANEJAMENTO", BRANCO),
                                      ("E RESULTADOS", VERDE))):
        d.texto(s, 0.88, 2.72 + 0.72 * i, 8.6, 0.72, linha, tam=42, negrito=True,
                cor=cor, ancora="ctr")
    d.texto(s, 0.92, 5.42, 9.5, 0.40,
            f"Resultados de {r['rotulo_fechado']}  ·  Planejamento de {r['rotulo_novo']}",
            tam=14, cor=VERDE_CLARO_CAPA, ancora="ctr")
    d.texto(s, 0.92, 5.85, 8.6, 0.38, "TIME COMERCIAL  ·  SDR  ·  EV  ·  EC",
            tam=11, negrito=True, cor=LARANJA_CLARO, espaco=250, ancora="ctr")
    s.notes_slide.notes_text_frame.text = (
        f"RPeR {r['rotulo_novo']}. Gerado pelo HIPO em "
        f"{r['gerado_em'].astimezone(FUSO_OPERACAO):%d/%m/%Y %H:%M}. "
        f"Resultados de {r['rotulo_fechado']}."
    )


def _secao(d: _Desenho, numero: int, squad: str, pessoas: list[dict]):
    s = d.slide()
    d.imagem(s, "bg_secao.png", 0, 0, 13.333, 7.5)
    d.imagem(s, "swoosh_laranja.png", -2.00, 4.70, 12.00, 1.98)
    d.imagem(s, "swoosh_verde.png", 4.00, 5.40, 12.00, 1.98)
    d.texto(s, 0.90, 1.75, 3.0, 1.50, f"{numero:02d}", tam=96, negrito=True, cor=BRANCO,
            ancora="b")
    d.texto(s, 0.98, 3.15, 8.0, 0.40, f"BLOCO {numero + 1:02d}", tam=13, negrito=True,
            cor=LARANJA_CLARO, espaco=300, ancora="ctr")
    d.texto(s, 0.90, 3.50, 11.0, 1.10, regras.NOME_SQUAD[squad], tam=40, negrito=True,
            cor=BRANCO, ancora="ctr")
    nomes = "  ·  ".join(p["nome"] for p in pessoas) or f"Nenhuma pessoa ativa com o cargo {squad}"
    d.texto(s, 0.95, 4.62, 11.0, 0.50, nomes, tam=12.5, cor=VERDE_CLARO_SECAO, ancora="ctr")
    d.cartao(s, 10.31, 0.32, 2.62, 0.84, raio=0.2)
    d.imagem(s, "logo.png", 10.47, 0.46, 2.30, 0.53)


def _estilo_resultado(linhas_calc):
    def estilo(i, j):
        linha = linhas_calc[i - 1]
        if j == 1:
            return {"fundo": META_FUNDO, "cor": LARANJA}
        if j == 3:
            return {"fundo": META_FUNDO,
                    "cor": COR_ATINGIMENTO.get(linha["carinha"], TEXTO_2)}
        if j == 2:
            return {"cor": VERDE_ESCURO}
        return None
    return estilo


def _nota_posicao(squad, r):
    posicoes = [i.rotulo for i in regras.INDICADORES[squad] if i.posicao]
    gerado = r["gerado_em"].astimezone(FUSO_OPERACAO)
    base = f"Gerado pelo HIPO em {gerado:%d/%m/%Y às %H:%M}."
    if posicoes:
        base += (" " + ", ".join(p.title() for p in posicoes)
                 + f": posição no momento da geração. Demais: fluxo de {r['rotulo_fechado']}.")
    return base


def _resultados(d: _Desenho, r: dict, squad: str, textos: dict):
    sq = r["squads"][squad]
    s = d.slide()
    d.fundo_conteudo(s, f"{squad} — RESULTADOS ", _mes_maiusculo(r["ano"], r["mes"]),
                     f"Realizado de {r['rotulo_fechado']} contra a meta do squad")
    linhas = [["INDICADOR", "META", "REALIZADO", "ATING. %"]] + [
        [l["rotulo"], l["meta_txt"], l["realizado_txt"], l["atingimento_txt"]]
        for l in sq["total"]
    ]
    alt = min(0.42, 4.95 / len(linhas))
    d.cartao(s, 0.50, 1.50, 7.54, alt * len(linhas) + 0.24, raio=0.04)
    d.tabela(s, 0.62, 1.62, [3.10, 1.40, 1.50, 1.30], linhas, altura_linha=alt,
             estilos=_estilo_resultado(sq["total"]))

    # Quatro quadros grandes, 2x2.
    for k, chave in enumerate(KPIS[squad]):
        x = 8.35 if k % 2 == 0 else 10.63
        y = 1.50 if k < 2 else 3.10
        valor = regras.texto(sq["total"], chave)
        d.kpi(s, x, y, 2.14, 1.45, valor, ROTULO_KPI[(squad, chave)])

    d.cartao(s, 8.35, 4.72, 4.42, 1.95, fundo=BEGE, raio=0.08)
    d.rotulo_cartao(s, 8.58, 4.84, 4.0, "LEITURA DO MÊS")
    d.texto(s, 8.58, 5.16, 4.0, 1.46, textos[squad]["leitura"], tam=10.5, cor=TEXTO)
    d.rodape(s, _nota_posicao(squad, r))


def _por_pessoa(d: _Desenho, r: dict, squad: str, textos: dict):
    sq = r["squads"][squad]
    pessoas = sq["pessoas"]
    s = d.slide()
    d.fundo_conteudo(s, NOME_POR_PESSOA[squad], "",
                     f"{r['rotulo_fechado'].capitalize()} · realizado e, abaixo, a meta individual")

    cab = ["INDICADOR"] + [p["rotulo"] for p in pessoas] + ["SQUAD"]
    corpo = []
    for k, ind in enumerate(regras.INDICADORES[squad]):
        linha = [ind.rotulo]
        for bloco in [p["indicadores"] for p in pessoas] + [sq["total"]]:
            l = bloco[k]
            cel = [[(l["realizado_txt"], {})]]
            if l["meta_txt"]:
                cor = COR_ATINGIMENTO.get(l["carinha"], TEXTO_2)
                cel.append([(f"meta {l['meta_txt']}", {"tam": 7.5, "negrito": False,
                                                        "cor": TEXTO_2}),
                            (f"  {l['atingimento_txt']}", {"tam": 7.5, "cor": cor})])
            linha.append(cel)
        corpo.append(linha)
    n_val = len(pessoas) + 1
    larg_val = min(1.45, 4.80 / n_val)
    larguras = [7.60 - larg_val * n_val] + [larg_val] * n_val
    linhas = [cab] + corpo
    alt = min(0.46, 4.80 / len(linhas))
    ultima = len(larguras) - 1

    def estilo(i, j):
        if j == ultima:
            return {"fundo": META_FUNDO, "cor": VERDE_ESCURO}
        return {"tam": 10 if j == 0 else 11}

    d.cartao(s, 0.50, 1.48, 7.84, alt * len(linhas) + 0.24, raio=0.04)
    d.tabela(s, 0.62, 1.60, larguras, linhas, altura_linha=alt, estilos=estilo, zebra=True)

    # Um cartão de comentário por pessoa, empilhados na coluna da direita.
    if not pessoas:
        d.cartao(s, 8.70, 1.60, 4.05, 1.2, raio=0.08)
        d.texto(s, 8.94, 1.72, 3.6, 0.9, f"Nenhuma pessoa ativa com o cargo {squad}.",
                tam=11, cor=TEXTO_2, ancora="ctr")
    else:
        n = len(pessoas)
        disponivel, vao = 4.80, 0.14
        h = (disponivel - vao * (n - 1)) / n
        tam = 10.5 if n <= 2 else (9.5 if n == 3 else 8.5)
        for k, p in enumerate(pessoas):
            y = 1.60 + k * (h + vao)
            d.cartao(s, 8.70, y, 4.05, h, raio=0.06)
            d.rotulo_cartao(s, 8.94, y + 0.10, 3.6, p["rotulo"], cor=VERDE_ESCURO)
            d.texto(s, 8.94, y + 0.42, 3.6, h - 0.50,
                    textos[squad]["pessoas"].get(str(p["id"]), ""), tam=tam, cor=TEXTO)


def _pipeline(d: _Desenho, r: dict, textos: dict):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION

    ev = r["squads"]["EV"]
    g = ev["graficos"]
    s = d.slide()
    d.fundo_conteudo(s, "PIPELINE POR ", "EXECUTIVO",
                     f"Oportunidades ativas na posição da geração · vendas de {r['rotulo_fechado']}")
    d.texto(s, 0.60, 1.36, 6.0, 0.30, "OPORTUNIDADES POR FASE", tam=11, negrito=True,
            cor=LARANJA_CLARO, espaco=100)
    d.texto(s, 7.10, 1.36, 6.0, 0.30, "PIPELINE EM TICKET MENSAL (R$)", tam=11,
            negrito=True, cor=LARANJA_CLARO, espaco=100)

    nomes = g["nomes"] or ["—"]
    zeros = [0] * len(nomes)

    def estilo_eixos(chart):
        chart.font.name = FONTE
        chart.font.size = d.Pt(10)
        chart.font.color.rgb = d._cor(BRANCO)
        cat = chart.category_axis
        cat.tick_labels.font.color.rgb = d._cor(BRANCO)
        cat.tick_labels.font.bold = True
        cat.format.line.color.rgb = d._cor("6B7F66")
        cat.has_major_gridlines = False
        val = chart.value_axis
        val.visible = False
        val.has_major_gridlines = True
        val.major_gridlines.format.line.color.rgb = d._cor("2A4426")

    fases = CategoryChartData()
    fases.categories = nomes
    fases.add_series("Em negociação", g["em_negociacao"] or zeros)
    fases.add_series(f"Vendas em {regras.rotulo_mes(r['ano'], r['mes'], com_ano=False)}",
                     g["concluidas"] or zeros)
    fases.add_series("Outras fases", g["outras_fases"] or zeros)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, d.In(0.50), d.In(1.72),
                            d.In(6.30), d.In(4.02), fases)
    c1 = gf.chart
    estilo_eixos(c1)
    c1.has_title = False
    c1.has_legend = True
    c1.legend.position = XL_LEGEND_POSITION.BOTTOM
    c1.legend.include_in_layout = False
    c1.legend.font.color.rgb = d._cor(BRANCO)
    c1.legend.font.size = d.Pt(10)
    plot = c1.plots[0]
    plot.gap_width = 90
    plot.overlap = 100
    plot.has_data_labels = True
    plot.data_labels.position = XL_LABEL_POSITION.CENTER
    plot.data_labels.font.color.rgb = d._cor(BRANCO)
    plot.data_labels.font.bold = True
    plot.data_labels.number_format = '0;;;'
    plot.data_labels.number_format_is_linked = False
    for serie, cor in zip(plot.series, (LARANJA, VERDE, CINZA_GRAFICO)):
        serie.format.fill.solid()
        serie.format.fill.fore_color.rgb = d._cor(cor)

    pip = CategoryChartData(number_format='"R$" #,##0')
    pip.categories = nomes
    pip.add_series("Pipeline (R$)", g["pipeline"] or zeros)
    gp = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, d.In(7.00), d.In(1.72),
                            d.In(5.85), d.In(4.02), pip)
    c2 = gp.chart
    estilo_eixos(c2)
    c2.has_legend = False
    c2.has_title = False
    p2 = c2.plots[0]
    p2.gap_width = 90
    p2.vary_by_categories = False
    p2.has_data_labels = True
    p2.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    p2.data_labels.font.color.rgb = d._cor(BRANCO)
    p2.data_labels.font.bold = True
    p2.data_labels.font.size = d.Pt(12)
    p2.data_labels.number_format = '"R$" #,##0'
    p2.data_labels.number_format_is_linked = False
    serie = p2.series[0]
    serie.format.fill.solid()
    serie.format.fill.fore_color.rgb = d._cor(VERDE)
    # Barras alternam verde e laranja, como no original.
    for k in range(len(nomes)):
        if k % 2 == 1:
            ponto = serie.points[k]
            ponto.format.fill.solid()
            ponto.format.fill.fore_color.rgb = d._cor(LARANJA)

    d.cartao(s, 0.55, 5.86, 12.25, 0.90, fundo=BEGE, raio=0.1)
    d.texto(s, 0.85, 5.94, 11.65, 0.74, textos["EV"]["pipeline"], tam=11, cor=TEXTO,
            ancora="ctr")


def _top(d: _Desenho, r: dict, textos: dict):
    ev = r["squads"]["EV"]
    neg = ev["negociacoes"]
    s = d.slide()
    d.fundo_conteudo(s, "TOP OPORTUNIDADES ", "EM NEGOCIAÇÃO",
                     "Maiores tickets mensais abertos na fase Negociação, na posição da geração")
    cab = ["#", "CONTA", "EXECUTIVO", "ÚLT. FUP", "TICKET MENSAL"]
    corpo = [
        [str(k + 1), (o["empresa"] or "—")[:42], o["executivo"],
         regras.data_curta(o["ultimo_fup"], FUSO_OPERACAO),
         regras.formatar(o["valor"], "moeda")]
        for k, o in enumerate(neg["top"])
    ] or [["—", "Nenhuma negociação aberta com EV envolvido", "", "", ""]]
    linhas = [cab] + corpo
    alt = 0.44 if len(linhas) <= 11 else 0.40

    def estilo(i, j):
        if j == 0:
            return {"fundo": LARANJA if i == 1 else VERDE, "cor": BRANCO, "alinhar": "ctr",
                    "tam": 10}
        if j == 1:
            return {"tam": 10, "alinhar": "l"}
        if j in (2, 3):
            return {"tam": 10, "cor": TEXTO_2, "negrito": False}
        return {"tam": 10.5, "cor": VERDE_ESCURO}

    d.cartao(s, 0.50, 1.38, 9.24, alt * len(linhas) + 0.24, raio=0.03)
    d.tabela(s, 0.62, 1.50, [0.50, 4.30, 1.60, 1.10, 1.50], linhas, altura_linha=alt,
             estilos=estilo, zebra=True)

    n = len(neg["top"])
    d.kpi(s, 10.05, 1.50, 2.72, 1.50, regras.formatar(neg["soma_top"], "moeda"),
          f"SOMA DO\nTOP {n}" if n else "SOMA DO TOP")
    d.kpi(s, 10.05, 3.16, 2.72, 1.50, regras.formatar(neg["abertas"], "inteiro"),
          "NEGOCIAÇÕES\nABERTAS")
    d.cartao(s, 10.05, 4.82, 2.72, 1.80, fundo=BEGE, raio=0.08)
    d.rotulo_cartao(s, 10.28, 4.94, 2.3, "FOCO")
    d.texto(s, 10.28, 5.26, 2.36, 1.30, textos["EV"]["foco"], tam=10, cor=TEXTO)
    d.rodape(s, "Ticket mensal = mensalidade da oportunidade no HIPO. Últ. FUP = última "
                "tarefa concluída da oportunidade. " + _nota_posicao("EV", r).split(".")[0] + ".")


def _planejamento(d: _Desenho, r: dict, squad: str, textos: dict):
    sq = r["squads"][squad]
    pl = sq["planejamento"]
    mes_novo = _mes_maiusculo(r["ano_novo"], r["mes_novo"])
    s = d.slide()
    d.fundo_conteudo(s, f"PLANEJAMENTO {mes_novo} — ", squad,
                     "Metas lançadas no HIPO · células vazias são definidas na reunião")

    linhas = [[f"{squad} — SQUAD", f"META {mes_novo}"]] + [
        [t["rotulo"], t["meta_txt"]] for t in pl["time"]
    ]
    alt = min(0.40, 3.55 / len(linhas))
    altura = alt * len(linhas) + 0.24
    d.cartao(s, 0.50, 1.54, 6.24, altura, raio=0.04)
    d.tabela(s, 0.62, 1.66, [4.00, 2.00], linhas, altura_linha=alt,
             estilos=lambda i, j: {"fundo": META_FUNDO, "cor": LARANJA} if j == 1 else None)

    y_base = 1.54 + altura + 0.18
    d.cartao(s, 0.62, y_base, 6.00, 6.62 - y_base, raio=0.08)
    d.rotulo_cartao(s, 0.88, y_base + 0.12, 5.5, f"BASE DE PARTIDA ({_mes_maiusculo(r['ano'], r['mes'])})",
                    cor=VERDE_ESCURO)
    d.texto(s, 0.88, y_base + 0.46, 5.5, 6.62 - y_base - 0.56,
            [regras.base_de_partida(squad, sq),
             [("Use estes números como piso ao definir a meta.", {"tam": 9.5, "cor": TEXTO_2,
                                                                  "negrito": False})]],
            tam=11, negrito=True, cor=TEXTO, entre=4)

    # Metas por pessoa.
    cab = [f"META POR {squad}"] + pl["colunas"]
    corpo = [[p["rotulo"]] + p["metas"] for p in pl["pessoas"]] or [["—", "", ""]]
    linhas_p = [cab] + corpo
    alt_p = min(0.40, 2.0 / len(linhas_p))
    altura_p = alt_p * len(linhas_p) + 0.24
    d.cartao(s, 6.98, 1.54, 5.84, altura_p, raio=0.05)
    d.tabela(s, 7.10, 1.66, [2.00, 1.80, 1.80], linhas_p, altura_linha=alt_p,
             estilos=lambda i, j: {"fundo": META_FUNDO, "cor": LARANJA} if j > 0 else None)

    y_acoes = 1.54 + altura_p + 0.18
    d.cartao(s, 7.10, y_acoes, 5.60, 6.62 - y_acoes, fundo=BEGE, raio=0.06)
    d.rotulo_cartao(s, 7.36, y_acoes + 0.14, 5.0, "AÇÕES DO MÊS")
    d.marcadores(s, 7.40, y_acoes + 0.50, 5.10, 6.62 - y_acoes - 0.60,
                 textos[squad]["acoes"], tam=10.5)


def montar_pptx(r: dict, textos: dict) -> bytes:
    """O RPeR inteiro, em memória. Não grava em disco (mesma razão da proposta)."""
    d = _Desenho()
    _capa(d, r)
    for numero, squad in enumerate(regras.SQUADS, start=1):
        _secao(d, numero, squad, r["squads"][squad]["pessoas"])
        _resultados(d, r, squad, textos)
        _por_pessoa(d, r, squad, textos)
        if squad == "EV":
            _pipeline(d, r, textos)
            _top(d, r, textos)
        _planejamento(d, r, squad, textos)
    buffer = BytesIO()
    d.prs.save(buffer)
    return buffer.getvalue()


def nome_do_arquivo(r: dict, formato: str) -> str:
    """RPeR_OUTUBRO_2026_CONTROLLER_MEDSEG.pptx — o nome que a operação já usava."""
    mes = regras.rotulo_mes(r["ano_novo"], r["mes_novo"], com_ano=False).upper()
    mes = (mes.replace("Ç", "C").replace("Ã", "A"))
    return f"RPeR_{mes}_{r['ano_novo']}_CONTROLLER_MEDSEG.{formato}"
